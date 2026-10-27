"""Scrittore DOCX dall'albero markup (python-docx, stili Word nativi).

Titoli → "Heading N", citazioni → "Quote", elenchi → "List Bullet/Number N" con la
numerazione che riparte a ogni elenco, codice → paragrafo monospazio ombreggiato,
tabelle → "Table Grid" con le celle unite, immagini incorporate (mai scaricate).
"""

from __future__ import annotations

from io import BytesIO

from ..docmodel.writers import _apply_hyperlink, _png_from, _restarted_num_id
from .common import WriteReport, as_blocks, heading_level, segments, table_grid
from .tree import HEADINGS, MarkupDocument, Node

_MONO = "Consolas"
_MAX_WIDTH_IN = 6.0


def _set_mono(run) -> None:
    from docx.oxml.ns import qn

    run.font.name = _MONO
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), _MONO)


def _picture_width(data: bytes):
    from docx.shared import Inches

    try:
        from PIL import Image

        with Image.open(BytesIO(data)) as image:
            dpi = image.info.get("dpi", (96, 96))[0] or 96
            width_in = image.width / float(dpi)
    except Exception:
        return Inches(_MAX_WIDTH_IN)
    return Inches(min(width_in, _MAX_WIDTH_IN))


def _picture_data(document: MarkupDocument, src: str) -> bytes | None:
    found = document.images.get(src)
    if not found:
        return None
    data, ext = found
    if ext in ("png", "jpg", "gif", "bmp", "tif"):
        return data
    return _png_from(data)          # webp e simili: Word li vuole in PNG


class _DocxWriter:
    def __init__(self, document: MarkupDocument, report: WriteReport) -> None:
        from docx import Document as DocxDocument

        self.document = document
        self.report = report
        self.docx = DocxDocument()
        if document.title:
            self.docx.core_properties.title = document.title

    def paragraph(self, style: str | None = None):
        if style:
            try:
                return self.docx.add_paragraph(style=style)
            except KeyError:
                pass
        return self.docx.add_paragraph()

    def runs(self, paragraph, children, *, mono: bool = False, bold: bool = False) -> None:
        for seg in segments(children):
            if seg.kind == "br":
                paragraph.add_run().add_break()
                continue
            if seg.kind == "img":
                data = _picture_data(self.document, seg.src)
                if data is None:
                    self.report.dropped_images += 1
                    continue
                try:
                    paragraph.add_run().add_picture(BytesIO(data), width=_picture_width(data))
                except Exception:
                    self.report.dropped_images += 1
                continue
            run = paragraph.add_run(seg.text)
            marks = seg.marks
            if "strong" in marks or bold:
                run.font.bold = True
            if "em" in marks:
                run.font.italic = True
            if "u" in marks:
                run.font.underline = True
            if "s" in marks:
                run.font.strike = True
            if "sup" in marks:
                run.font.superscript = True
            if "sub" in marks:
                run.font.subscript = True
            if "code" in marks or mono:
                _set_mono(run)
            if seg.href:
                _apply_hyperlink(paragraph, run, seg.href)

    def blocks(self, nodes: list[Node], *, quote: bool = False, indent_pt: float = 0.0) -> None:
        for node in nodes:
            self.block(node, quote=quote, indent_pt=indent_pt)

    def block(self, node: Node, *, quote: bool = False, indent_pt: float = 0.0) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        tag = node.tag
        if tag in HEADINGS:
            self.runs(self.paragraph(f"Heading {heading_level(tag)}"), node.children)
        elif tag in ("p", "dd", "dt", "figcaption"):
            style = "Quote" if quote else ("Caption" if tag == "figcaption" else None)
            paragraph = self.paragraph(style)
            self.runs(paragraph, node.children, bold=(tag == "dt"))
            if indent_pt or tag == "dd":
                paragraph.paragraph_format.left_indent = Pt(indent_pt + (18 if tag == "dd" else 0))
            if len(node.children) == 1 and isinstance(node.children[0], Node) and node.children[0].tag == "img":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tag == "pre":
            paragraph = self.paragraph()
            lines = node.text().rstrip("\n").split("\n")
            for index, line in enumerate(lines):
                run = paragraph.add_run(line)
                _set_mono(run)
                run.font.size = Pt(9.5)
                if index < len(lines) - 1:
                    run.add_break()
            _shade(paragraph)
            if indent_pt:
                paragraph.paragraph_format.left_indent = Pt(indent_pt)
        elif tag == "blockquote":
            self.blocks(as_blocks(node), quote=True, indent_pt=indent_pt)
        elif tag in ("ul", "ol"):
            self.list(node, level=1, indent_pt=indent_pt)
        elif tag == "hr":
            _rule(self.paragraph())
        elif tag == "table":
            self.table(node)
        elif tag in ("figure", "dl", "li", "td", "th"):
            self.blocks(as_blocks(node), quote=quote, indent_pt=indent_pt)
        elif tag == "img":
            paragraph = self.paragraph()
            self.runs(paragraph, [node])

    def list(self, node: Node, *, level: int, indent_pt: float = 0.0) -> None:
        from docx.shared import Pt

        ordered = node.tag == "ol"
        style_level = min(level, 3)
        base = "List Number" if ordered else "List Bullet"
        style = base if style_level == 1 else f"{base} {style_level}"
        restarted: dict[str, int] = {}
        start = int(node.attrs.get("start", "1") or 1)
        for item in node.children:
            if not isinstance(item, Node):
                continue
            first = True
            for block in as_blocks(item):
                if block.tag in ("ul", "ol"):
                    self.list(block, level=level + 1, indent_pt=indent_pt)
                    continue
                if first and block.tag in ("p",) + HEADINGS:
                    paragraph = self.paragraph(style)
                    if ordered:
                        num_id = _restarted_num_id(self.docx, style, restarted, start)
                        if num_id is not None:
                            num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                            num_pr.get_or_add_numId().val = num_id
                    self.runs(paragraph, block.children)
                    if level > 3:
                        paragraph.paragraph_format.left_indent = Pt(indent_pt + 18 * level)
                    first = False
                else:
                    # Paragrafi successivi della stessa voce: rientrati sotto il testo.
                    self.block(block, indent_pt=indent_pt + 18 * level)
                    first = False

    def table(self, node: Node) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        grid, header_rows, _merged = table_grid(node)
        if not grid or not grid[0]:
            return
        rows, cols = len(grid), len(grid[0])
        table = self.docx.add_table(rows=rows, cols=cols)
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        align = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for r, line in enumerate(grid):
            for c, cell in enumerate(line):
                if cell is None:
                    continue
                target = table.cell(r, c)
                colspan = int(cell.attrs.get("colspan", "1") or 1)
                rowspan = int(cell.attrs.get("rowspan", "1") or 1)
                if colspan > 1 or rowspan > 1:
                    target = target.merge(table.cell(min(rows - 1, r + rowspan - 1),
                                                     min(cols - 1, c + colspan - 1)))
                header = r < header_rows or cell.tag == "th"
                first = True
                for block in as_blocks(cell):
                    paragraph = target.paragraphs[0] if first else target.add_paragraph()
                    first = False
                    if block.tag in ("ul", "ol", "table", "pre", "blockquote"):
                        self.report.note("Lists, code and nested tables inside table cells became plain text.")
                        paragraph.add_run(" ".join(block.text().split()))
                    else:
                        self.runs(paragraph, block.children, bold=header)
                    if cell.attrs.get("align") in align:
                        paragraph.alignment = align[cell.attrs["align"]]
        self.docx.add_paragraph()      # Word non accetta una tabella in coda al corpo


def _shade(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F2F2F2")
    paragraph._p.get_or_add_pPr().append(shading)


def _rule(paragraph) -> None:
    """Linea orizzontale: bordo inferiore del paragrafo (come il "---" di Word)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def write_docx(document: MarkupDocument, path) -> WriteReport:
    report = WriteReport()
    writer = _DocxWriter(document, report)
    writer.blocks([c for c in document.body.children if isinstance(c, Node)])
    writer.docx.save(str(path))
    return report
