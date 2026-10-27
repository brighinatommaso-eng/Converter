"""Lettore DOCX → IR (python-docx, MIT).

Serve alle coppie docx→md/txt: pandoc (usato un tempo) produceva "grid table" e
appiattiva i titoli nel front-matter. Qui la struttura Word (Heading N, liste,
tabelle con header) arriva all'IR e quindi a un Markdown GFM pulito.

Il lettore conserva anche ciò che Word mette fuori dai semplici `paragraph.runs`:
hyperlink, immagini inline, testatine/piè di pagina e merge verticali/orizzontali
delle celle. Sono tutti dati dell'IR, non scorciatoie del singolo writer.
"""

from __future__ import annotations

from pathlib import Path

from .types import Document, ImageBlock, Line, ListBlock, ListItem, Paragraph, Run, Table, TableCell

_IMAGE_EXT = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/gif": "gif",
    "image/bmp": "bmp",
    "image/tiff": "tif",
    "image/webp": "webp",
    "image/svg+xml": "svg",
}


def _run_style(run, style) -> tuple[bool, bool, str, float, str]:
    """(bold, italic, family, size_pt, #RRGGBB) risolvendo l'eredità dallo stile."""
    bold = run.bold if run.bold is not None else bool(style and style.font.bold)
    italic = run.italic if run.italic is not None else bool(style and style.font.italic)
    family = run.font.name or (style.font.name if style and style.font.name else "")
    size = 0.0
    if run.font.size is not None:
        size = float(run.font.size.pt)
    elif style is not None and style.font.size is not None:
        size = float(style.font.size.pt)
    color = ""
    try:
        rgb = run.font.color.rgb
        if rgb is not None:
            color = f"#{rgb}"
    except (AttributeError, ValueError):
        color = ""
    return bold, italic, family, size, color


# Contenitori "trasparenti" di un paragrafo: il loro testo è testo del corpo.
# w:ins = revisione inserita, w:fldSimple = risultato di un campo, w:sdt =
# controllo contenuto, w:smartTag/w:customXml = marcature. w:del e w:moveFrom
# (testo cancellato/spostato via) restano fuori, come in Word a revisioni accettate.
_TRANSPARENT = {"ins", "moveTo", "smartTag", "customXml", "fldSimple", "sdt", "sdtContent",
                "dir", "bdo"}


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _flatten(element, wanted: str):
    """Figli `wanted` di `element`, attraversando sdt/customXml/ins (righe e celle
    dentro un controllo contenuto venivano saltate)."""
    for child in element.iterchildren():
        tag = _local(child.tag)
        if tag == wanted:
            yield child
        elif tag in _TRANSPARENT:
            yield from _flatten(child, wanted)


def _iter_runs(docx_paragraph):
    """(run, URL) nell'ordine del documento, ovunque si trovino nel paragrafo.

    `paragraph.runs` (e `iter_inner_content`) vedono solo i `w:r` diretti e gli
    hyperlink: il testo di una revisione inserita, del risultato di un campo o
    di un controllo contenuto spariva in Markdown e in TXT senza avviso.
    """
    from docx.text.hyperlink import Hyperlink
    from docx.text.run import Run as DocxRun

    def walk(element, link: str):
        for child in element.iterchildren():
            tag = _local(child.tag)
            if tag == "r":
                yield DocxRun(child, docx_paragraph), link
            elif tag == "hyperlink":
                try:
                    url = Hyperlink(child, docx_paragraph).url or ""
                except Exception:
                    url = ""       # relazione rotta: il testo resta, il link no
                yield from walk(child, url or link)
            elif tag in _TRANSPARENT:
                yield from walk(child, link)

    yield from walk(docx_paragraph._p, "")


def _textbox_paragraphs(run, docx) -> list[Paragraph]:
    """Paragrafi delle caselle di testo ancorate al run (w:txbxContent).

    Word salva la casella due volte (mc:Choice moderna + mc:Fallback VML): si
    legge solo la prima, altrimenti il testo comparirebbe doppio."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph as DocxParagraph

    fallback = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
    paragraphs: list[Paragraph] = []
    for box in run._r.iter(qn("w:txbxContent")):
        if any(ancestor.tag == fallback for ancestor in box.iterancestors()):
            continue
        for child in _flatten(box, "p"):
            paragraph, _, _ = _paragraph(DocxParagraph(child, docx))
            if paragraph.text.strip():
                paragraphs.append(paragraph)
    return paragraphs


def _run_ir(run, style, link: str = "") -> Run | None:
    text = run.text
    if not text:
        return None
    bold, italic, family, size, color = _run_style(run, style)
    return Run(text=text, bold=bold, italic=italic, font=family, size=size,
               color=color, link=link)


def _paragraph_from_runs(runs: list[Run], style_name: str) -> Paragraph:
    low = style_name.lower()
    if low == "title":
        return Paragraph(lines=[Line(runs=runs)], style="title", level=1)
    if low.startswith("heading "):
        try:
            level = int(low.split()[-1])
        except ValueError:
            level = 1
        return Paragraph(lines=[Line(runs=runs)], style="heading", level=level)
    return Paragraph(lines=[Line(runs=runs)])


def _paragraph(docx_paragraph) -> tuple[Paragraph, str, int]:
    """(paragrafo IR, nome stile, livello) da un paragrafo Word."""
    style = docx_paragraph.style
    style_name = style.name if style is not None else ""
    runs = [run for item, link in _iter_runs(docx_paragraph)
            if (run := _run_ir(item, style, link)) is not None]
    paragraph = _paragraph_from_runs(runs, style_name)
    return paragraph, style_name, paragraph.level


def _image_extension(part) -> str:
    content_type = str(getattr(part, "content_type", "") or "").lower()
    if content_type in _IMAGE_EXT:
        return _IMAGE_EXT[content_type]
    suffix = str(getattr(getattr(part, "partname", None), "ext", "") or "").lower()
    if suffix:
        return "jpg" if suffix == "jpeg" else suffix
    return "png"


def _drawing_size(drawing) -> tuple[float, float]:
    from docx.oxml.ns import qn

    for extent in drawing.iter(qn("wp:extent")):
        try:
            return float(int(extent.get("cx"))) / 12700.0, float(int(extent.get("cy"))) / 12700.0
        except (TypeError, ValueError):
            break
    return 0.0, 0.0


def _images_from_run(run) -> list[ImageBlock]:
    """Immagini inline/ancorate contenute in un run Word."""
    from docx.oxml.ns import qn

    images: list[ImageBlock] = []
    for drawing in run._r.iter(qn("w:drawing")):
        for blip in drawing.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if not rid:
                continue
            try:
                part = run.part.related_parts[rid]
                data = part.blob
            except (KeyError, AttributeError):
                continue
            width, height = _drawing_size(drawing)
            images.append(ImageBlock(data=data, ext=_image_extension(part),
                                     width_pt=width, height_pt=height))
    return images


def _paragraph_parts(docx_paragraph) -> tuple[list[Paragraph | ImageBlock], str, list[Paragraph]]:
    """Blocchi di un paragrafo, spezzati attorno alle immagini.

    Un'immagine inline non è testo: diventa un `ImageBlock` nello stesso punto
    dell'ordine di lettura. Se il paragrafo ha solo l'immagine non viene creato
    un paragrafo vuoto.
    """
    style = docx_paragraph.style
    style_name = style.name if style is not None else ""
    runs: list[Run] = []
    blocks: list[Paragraph | ImageBlock] = []

    def flush() -> None:
        if any(run.text.strip() for run in runs):
            blocks.append(_paragraph_from_runs(list(runs), style_name))
        runs.clear()

    boxes: list[Paragraph] = []
    for item, link in _iter_runs(docx_paragraph):
        run = _run_ir(item, style, link)
        if run is not None:
            runs.append(run)
        for image in _images_from_run(item):
            flush()
            blocks.append(image)
        boxes.extend(_textbox_paragraphs(item, docx_paragraph._parent))
    flush()
    # Il testo delle caselle segue il paragrafo che le ancora (ordine di lettura):
    # torna a parte perché non eredita lo stile (titolo/elenco) del paragrafo.
    return blocks, style_name, boxes


def _cell_paragraphs(tc, docx) -> list[Paragraph]:
    """Paragrafi di una cella, leggendo il `w:tc` grezzo (merge inclusi).

    Una tabella annidata non ha posto nella griglia dell'IR: ogni sua riga
    diventa un paragrafo con le celle separate da tabulazione, così il testo
    non sparisce (prima veniva ignorata del tutto)."""
    from docx.text.paragraph import Paragraph as DocxParagraph

    paragraphs = []
    for child in tc.iterchildren():
        tag = _local(child.tag)
        if tag in _TRANSPARENT:
            # Controllo contenuto/customXml attorno a paragrafi o tabelle della
            # cella: si scende (sdtPr e simili non sono p/tbl e vengono ignorati).
            paragraphs.extend(_cell_paragraphs(child, docx))
            continue
        if tag == "p":
            paragraph, _, _ = _paragraph(DocxParagraph(child, docx))
            if paragraph.text.strip():
                paragraphs.append(paragraph)
        elif tag == "tbl":
            for tr in _flatten(child, "tr"):
                runs: list[Run] = []
                for inner_tc in _flatten(tr, "tc"):
                    for inner in _cell_paragraphs(inner_tc, docx):
                        if runs:
                            runs.append(Run(text="\t"))
                        runs.extend(run for line in inner.lines for run in line.runs)
                if any(run.text.strip() for run in runs):
                    paragraphs.append(Paragraph(lines=[Line(runs=runs)]))
    return paragraphs


def _grid_span(tc) -> int:
    from docx.oxml.ns import qn

    tc_pr = tc.tcPr
    span = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
    try:
        return max(1, int(span.get(qn("w:val")))) if span is not None else 1
    except (TypeError, ValueError):
        return 1


def _vertical_merge(tc) -> str:
    """`restart`, `continue` o "" per una cella Word."""
    from docx.oxml.ns import qn

    tc_pr = tc.tcPr
    merge = tc_pr.find(qn("w:vMerge")) if tc_pr is not None else None
    if merge is None:
        return ""
    return merge.get(qn("w:val")) or "continue"


def _table_rows(table, docx) -> list[list[TableCell]]:
    """Righe con merge orizzontali e verticali rappresentati nell'IR.

    Una continuazione verticale non crea una cella nuova: incrementa `row_span`
    della cella di partenza. `col` tiene la posizione logica anche quando la
    cella coperta è stata omessa, così i writer possono ricostruire la griglia.
    """
    rows: list[list[TableCell]] = []
    active: dict[int, TableCell] = {}

    for tr in _flatten(table._tbl, "tr"):
        row: list[TableCell] = []
        continued: set[int] = set()
        started: set[int] = set()
        col = 0
        for tc in _flatten(tr, "tc"):
            span = _grid_span(tc)
            merge = _vertical_merge(tc)
            if merge == "continue":
                origin = active.get(col)
                if origin is not None:
                    origin.row_span += 1
                    continued.add(col)
                col += span
                continue

            for covered in range(col, col + span):
                active.pop(covered, None)
            cell = TableCell(paragraphs=_cell_paragraphs(tc, docx),
                             col=col, col_span=span)
            if merge == "restart":
                active[col] = cell
                started.add(col)
            row.append(cell)
            col += span

        # Un merge che non è continuato nella riga successiva è finito: non deve
        # restare appeso se il documento è malformato e salta una cella.
        for key in list(active):
            if key not in continued and key not in started:
                active.pop(key, None)
        rows.append(row)
    return rows


def _table_widths(table) -> list[float]:
    widths: list[float] = []
    try:
        for column in table.columns:
            try:
                widths.append(float(column.width.pt) if column.width else 0.0)
            except (AttributeError, ValueError):
                widths.append(0.0)
    except Exception:
        # Tabelle malformate senza w:tblGrid (Bug66263): la griglia testuale
        # resta valida, si perdono solo le larghezze.
        return []
    return [width for width in widths if width]


def _numbering_info(docx_paragraph) -> tuple[bool, int, int | None] | None:
    """(ordinato, livello, numId) dal `w:numPr`, risolvendo il formato in numbering.xml.

    Word non mette "List Bullet/Number" nello stile: la maggior parte delle liste
    reali è `List Paragraph` + `w:numPr`. Senza questo, ogni voce diventa un
    paragrafo e la struttura sparisce.
    """
    from docx.oxml.ns import qn

    p_pr = docx_paragraph._p.pPr
    num_pr = p_pr.numPr if p_pr is not None else None
    if num_pr is None or num_pr.numId is None:
        return None
    num_id = num_pr.numId.val
    if not num_id:                      # numId 0 = numerazione rimossa
        return None
    level = int(num_pr.ilvl.val) + 1 if num_pr.ilvl is not None else 1

    ordered = True
    try:
        numbering = docx_paragraph.part.numbering_part.element
    except (AttributeError, NotImplementedError):
        return ordered, level, num_id

    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(num_id):
            ref = num.find(qn("w:abstractNumId"))
            abstract_id = ref.get(qn("w:val")) if ref is not None else None
            break
    if abstract_id is None:
        return ordered, level, num_id

    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) != abstract_id:
            continue
        formats: dict[str, str] = {}
        for lvl in abstract.findall(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            formats[lvl.get(qn("w:ilvl"))] = fmt.get(qn("w:val")) if fmt is not None else ""
        # Un livello non dichiarato eredita il formato del primo: è quello che
        # fa Word quando il documento definisce solo il livello 0.
        fmt = formats.get(str(level - 1), formats.get("0", ""))
        if fmt == "bullet":
            ordered = False
        break
    return ordered, level, num_id


def _section_furniture(docx) -> tuple[list[Paragraph], list[Paragraph]]:
    """Testatine/piè di pagina delle sezioni, senza duplicare quelle collegate."""
    headers: list[Paragraph] = []
    footers: list[Paragraph] = []
    seen_headers: set[str] = set()
    seen_footers: set[str] = set()

    for section in docx.sections:
        for attr, target, seen in (("header", headers, seen_headers),
                                   ("footer", footers, seen_footers)):
            try:
                container = getattr(section, attr)
                if container.is_linked_to_previous:
                    continue
                paragraphs = container.paragraphs
            except Exception:
                continue
            for docx_paragraph in paragraphs:
                paragraph, _, _ = _paragraph(docx_paragraph)
                text = paragraph.text.strip()
                if text and text not in seen:
                    target.append(paragraph)
                    seen.add(text)
    return headers, footers


def read_docx(path: Path) -> Document:
    """DOCX → Document. Solleva ValueError con messaggio pulito se illeggibile."""
    from docx import Document as DocxDocument
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    try:
        docx = DocxDocument(str(path))
    except Exception as exc:
        raise ValueError(f"Could not read this DOCX: {exc}") from exc

    document = Document(source_format="docx")
    section = docx.sections[0] if docx.sections else None
    if section is not None:
        # Un DOCX può non dichiarare la sezione (Bug60341): None non è 0.
        if section.page_width is not None:
            document.page_width_pt = float(section.page_width.pt)
        if section.page_height is not None:
            document.page_height_pt = float(section.page_height.pt)
        if section.left_margin is not None:
            document.margin_left_pt = float(section.left_margin.pt)
        if section.right_margin is not None:
            document.margin_right_pt = float(section.right_margin.pt)

    document.headers, document.footers = _section_furniture(docx)

    pending_list: ListBlock | None = None

    def flush_list() -> None:
        nonlocal pending_list
        pending_list = None

    list_num_id: int | None = None   # numId del livello 1 dell'elenco aperto

    def body_children(element):
        """Figli del corpo, entrando nei controlli contenuto (w:sdt) a livello di
        blocco: prima il loro testo (moduli, modelli, campi) veniva saltato."""
        for child in element.iterchildren():
            tag = _local(child.tag)
            if tag in _TRANSPARENT:
                yield from body_children(child)
            elif tag in ("p", "tbl"):
                yield tag, child

    for tag, child in body_children(docx.element.body):
        if tag == "tbl":
            flush_list()
            table = DocxTable(child, docx)
            document.blocks.append(Table(rows=_table_rows(table, docx), borders=True,
                                         header_rows=1, widths_pt=_table_widths(table)))
            continue

        parts, style_name, boxes = _paragraph_parts(DocxParagraph(child, docx))
        if not parts and not boxes:
            continue
        low = style_name.lower()
        is_heading = low == "title" or low.startswith("heading ")

        info = None if is_heading else _numbering_info(DocxParagraph(child, docx))
        if info is None and not is_heading:
            is_bullet_style = low.startswith("list bullet")
            is_number_style = low.startswith("list number")
            if is_bullet_style or is_number_style:
                suffix = low.split()[-1]
                level = int(suffix) if suffix.isdigit() else 1
                info = (is_number_style, level, None)

        for part in parts:
            if isinstance(part, ImageBlock):
                flush_list()
                document.blocks.append(part)
                continue
            if is_heading:
                flush_list()
                document.blocks.append(part)
                continue
            if info is not None:
                ordered, level, num_id = info
                # Un elenco resta UNO anche se mescola voci numerate e puntate
                # (sotto-punti): ogni voce porta il proprio tipo. Un nuovo numId
                # al primo livello è invece un elenco Word distinto (riparte da 1).
                new_list = (pending_list is None
                            or (level == 1 and num_id is not None and list_num_id is not None
                                and num_id != list_num_id))
                if new_list:
                    pending_list = ListBlock(ordered=ordered)
                    document.blocks.append(pending_list)
                    list_num_id = None
                if level == 1 and num_id is not None:
                    list_num_id = num_id
                marker = "•" if not ordered else "1"
                pending_list.items.append(ListItem(marker=marker, level=level, paragraph=part,
                                                   ordered=ordered))
                continue
            flush_list()
            document.blocks.append(part)
        if boxes:
            flush_list()
            document.blocks.extend(boxes)

    if document.title == "":
        for block in document.blocks:
            if isinstance(block, Paragraph) and block.style == "title":
                document.title = block.text
                break
    return document
