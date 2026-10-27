"""Scrittore DOCX unico (modificabile + impaginato come l'originale)."""

from __future__ import annotations

from pathlib import Path

from converter.converters.docmodel.docx_faithful import _Band, _page_items, text_scale, word_font
from converter.converters.docmodel.types import (
    Document,
    ImageBlock,
    Line,
    PageBreak,
    Paragraph,
    Run,
)
from converter.converters.docmodel.writers import write_docx


def _line(text, x, y, *, size=10.0, font="Arial", width=None):
    width = width if width is not None else 5.0 * len(text)
    return Line(runs=[Run(text=text, size=size, font=font)], x=x, y=y, width=width, height=size * 1.2)


def _doc(blocks, size=(595.0, 842.0)):
    return Document(blocks=blocks, page_width_pt=size[0], page_height_pt=size[1],
                    source_format="pdf", page_sizes=[size] * (1 + sum(isinstance(b, PageBreak) for b in blocks)))


def test_missing_fonts_map_to_standard_ones_of_the_same_kind():
    assert word_font("MinionPro") == "Times New Roman"
    assert word_font("Helvetica Neue LT") == "Arial"
    assert word_font("Poppins") == "Arial"
    assert word_font("Courier Std") == "Courier New"
    assert word_font("Calibri") == "Calibri"            # installato: resta


def test_narrow_substituted_font_is_condensed_never_expanded():
    narrow = Paragraph(lines=[_line("The quick brown fox jumps over", 72, 700, font="MinionPro",
                                    width=120.0)])
    scale = text_scale(narrow)
    assert scale is not None and 66 <= scale < 98
    wide = Paragraph(lines=[_line("The quick brown fox jumps over", 72, 700, font="MinionPro",
                                  width=400.0)])
    assert text_scale(wide) == 98 or text_scale(wide) is None or text_scale(wide) <= 98
    real = Paragraph(lines=[_line("The quick brown fox", 72, 700, font="Arial")])
    assert text_scale(real) is None                     # font vero: nessuna correzione


def test_side_by_side_blocks_become_one_band_and_a_centered_title_does_not_merge_them():
    title = Paragraph(lines=[_line("Centered title", 200, 760, width=180)])
    left = [Paragraph(lines=[_line("left column", 72, 700 - 14 * i, width=200)]) for i in range(3)]
    right = [Paragraph(lines=[_line("right column", 320, 700 - 14 * i, width=200)]) for i in range(3)]
    items = _page_items([title, *left, *right], 72, 520)
    assert items[0] is title
    band = items[1]
    assert isinstance(band, _Band) and len(band.columns) == 2
    assert band.cells[0] == left and band.cells[1] == right


def test_pages_become_sections_with_their_size_and_columns_a_borderless_table(tmp_path: Path):
    import docx

    left = [Paragraph(lines=[_line("left text", 72, 700 - 14 * i, width=200)]) for i in range(3)]
    right = [Paragraph(lines=[_line("right text", 320, 700 - 14 * i, width=200)]) for i in range(3)]
    second = Paragraph(lines=[_line("Second page", 72, 500)])
    document = _doc([*left, *right, PageBreak(), second], size=(420.0, 595.0))
    out = tmp_path / "out.docx"
    write_docx(document, out)
    result = docx.Document(str(out))
    assert len(result.sections) == 2
    assert result.sections[1].page_width.pt == 420.0
    table = result.tables[0]
    assert [c.text.split("\n")[0] for c in table.rows[0].cells] == ["left text", "right text"]
    assert not table._tbl.xpath(".//w:tblBorders")      # invisibile


def test_shading_background_and_page_number_footer(tmp_path: Path):
    import docx

    boxed = Paragraph(lines=[_line("white on blue", 72, 700, font="Arial")], shading="#1E88E5")
    boxed.lines[0].runs[0].color = "#ffffff"
    background = ImageBlock(data=_png(), ext="png", width_pt=595, height_pt=842, x=0, y=0)
    document = _doc([background, boxed])
    document.footers = [Paragraph(lines=[_line("1", 290, 30)], role="page_number")]
    out = tmp_path / "out.docx"
    write_docx(document, out)
    xml = docx.Document(str(out)).element.xml
    assert 'w:fill="1E88E5"' in xml                       # riquadro colorato
    assert 'w:val="FFFFFF"' in xml                       # bianco conservato sopra lo sfondo
    assert 'behindDoc="1"' in xml                        # sfondo dietro al testo
    footer_xml = docx.Document(str(out)).sections[0].footer._element.xml
    assert "PAGE" in footer_xml                          # numero di pagina come campo


def _png() -> bytes:
    from io import BytesIO

    from PIL import Image

    buffer = BytesIO()
    Image.new("RGB", (8, 8), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()
