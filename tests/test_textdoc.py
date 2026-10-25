"""TextDocConverter: il motore strutturale (IR) per i documenti testuali.

Copre le cose che prima non esistevano o erano sbagliate:
  * pdf→md (la coppia non c'era proprio);
  * pdf→docx con STILI veri (Heading N, liste native, tabella) e non un clone;
  * pdf→txt con ordine di lettura e tabelle tab-separate (non colonne smontate);
  * docx→md con tabelle GFM invece delle grid table di pandoc.

La fixture "ricca" nasce da un docx costruito ad hoc e stampato in PDF da
LibreOffice (heading, paragrafo, tabella, lista, seconda pagina): è il PDF più
cattivo da ricostruire, quindi il test vale. Senza LibreOffice i test e2e si
saltano; la logica pura dell'IR è testata senza tool esterni.
"""

from pathlib import Path

import pytest

import converter.converters  # noqa: F401  (registra i converter)
from converter.converters import _ocr
from converter.converters.docmodel.pdfreader import _split_font, _xy_cut
from converter.converters.docmodel.types import Document, Line, Paragraph, Run
from converter.converters.textdoc import TextDocConverter
from converter.core import deps, engine, registry
from converter.core.job import ConversionJob
from converter.core.pipeline import run_job

REPORT = "Quarterly report 2026"


# --- PDF minimale scritto a mano (nessun tool esterno) ----------------------

def _write_text_pdf(path: Path, text: str, size: int = 18) -> None:
    """PDF con testo VERO (operatore Tj): Pillow sa scrivere solo PDF-immagine."""
    content = f"BT /F1 {size} Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode()
            + b" /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF\n")
    path.write_bytes(out)


def _job(src: Path, dst: str, out: Path, **options) -> ConversionJob:
    return ConversionJob(src_path=src, src_format=src.suffix.lstrip("."),
                         dst_format=dst, output_path=out, options=options)


# --- PDF "ricco": docx → PDF via LibreOffice --------------------------------

@pytest.fixture(scope="module")
def rich_pdf(tmp_path_factory) -> Path:
    if not deps.find_soffice():
        pytest.skip("LibreOffice required to build the rich PDF fixture")
    import subprocess

    from docx import Document as Docx

    folder = tmp_path_factory.mktemp("richpdf")
    doc = Docx()
    doc.add_heading(REPORT, 0)
    doc.add_heading("1. Summary", level=1)
    doc.add_paragraph("Revenue grew 12% over the previous quarter, driven by the services "
                      "business and by two large renewals. Costs stayed flat.")
    doc.add_paragraph("The second paragraph explains the methodology used to compute the "
                      "figures and the assumptions behind them.")
    doc.add_heading("2. Figures", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate([["Region", "Revenue", "Growth"],
                             ["North", "1,200", "8%"],
                             ["South", "940", "15%"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_heading("3. Next steps", level=2)
    for item in ("Expand the South region", "Hire two engineers", "Ship v2 in October"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()
    doc.add_heading("Appendix", level=1)
    doc.add_paragraph("A second page, to check page breaks and headers.")
    src = folder / "report.docx"
    doc.save(src)
    subprocess.run([deps.find_soffice(), "--headless", "--convert-to", "pdf",
                    "--outdir", str(folder), str(src)],
                   check=True, capture_output=True, env=deps.clean_subprocess_env())
    return folder / "report.pdf"


# --- Routing e disponibilità ------------------------------------------------

def test_routes_pdf_targets_to_structural_engine():
    engine.set_mode("FULL")
    for dst in ("docx", "odt", "rtf", "md", "txt", "html"):
        assert isinstance(registry.resolve("pdf", dst), TextDocConverter), dst


@pytest.mark.parametrize("mode", ["FULL", "PORTABLE"])
def test_pdf_docx_md_txt_available_in_both_modes(mode):
    """Non serve LibreOffice né Full mode: python-docx e pdfminer sono in-process."""
    engine.set_mode(mode)
    targets = {t.dst_format: t for t in registry.targets_for("pdf")}
    for dst in ("docx", "md", "txt", "html"):
        assert targets[dst].available is True, (dst, targets[dst].reason_unavailable)
        assert targets[dst].verdict.is_lossy, dst


def test_pdf_to_odt_rtf_need_libreoffice_and_full_mode(monkeypatch):
    monkeypatch.setattr(deps, "find_soffice", lambda: None)
    converter = TextDocConverter()
    assert converter.pair_available("pdf", "docx") == (True, "")
    for dst in ("odt", "rtf"):
        ok, reason = converter.pair_available("pdf", dst)
        assert ok is False and "LibreOffice" in reason

    monkeypatch.setattr(deps, "find_soffice", lambda: r"C:\soffice.exe")
    engine.set_mode("PORTABLE")
    ok, reason = converter.pair_available("pdf", "odt")
    assert ok is False and "Full mode" in reason


def test_available_requires_pdfminer_and_python_docx(monkeypatch):
    real = deps.has_module
    for missing in ("pdfminer", "docx"):
        monkeypatch.setattr(deps, "has_module",
                            lambda name, _m=missing: False if name == _m else real(name))
        status = TextDocConverter.available()
        assert status.ok is False and missing in status.reason + status.install_hint


def test_pdf_to_md_is_offered():
    """Non esisteva: `resolve("pdf","md")` tornava None."""
    targets = {t.dst_format for t in registry.targets_for("pdf")}
    assert "md" in targets
    assert isinstance(registry.resolve("pdf", "md"), TextDocConverter)


def test_pdf_to_txt_beats_the_flat_pdfminer_path():
    from converter.converters.pdf import PdfConverter

    assert isinstance(registry.resolve("pdf", "txt"), TextDocConverter)
    assert PdfConverter.priority < TextDocConverter.priority


def test_docx_markdown_targets_use_structural_engine():
    engine.set_mode("FULL")
    assert isinstance(registry.resolve("docx", "md"), TextDocConverter)
    assert isinstance(registry.resolve("docx", "txt"), TextDocConverter)
    # Le coppie office↔office e docx→html restano a LibreOffice (fedeltà colori).
    assert registry.resolve("docx", "docx").__class__.__name__ == "DocumentConverter"
    assert registry.resolve("docx", "html").__class__.__name__ == "DocumentConverter"


# --- validate(): scansioni e PDF rotti --------------------------------------

def _mock_ocr_chain(monkeypatch, *, present=True):
    monkeypatch.setattr(deps, "find_ocrmypdf", lambda: r"C:\ocrmypdf.exe" if present else None)
    monkeypatch.setattr(deps, "find_ghostscript", lambda: r"C:\gswin64c.exe" if present else None)
    monkeypatch.setattr(deps, "find_tesseract", lambda: r"C:\tesseract.exe" if present else None)


def test_validate_scanned_pdf_refused_without_ocr_chain(tmp_path, monkeypatch):
    from PIL import Image

    _mock_ocr_chain(monkeypatch, present=False)
    pdf = tmp_path / "scan.pdf"
    Image.new("RGB", (200, 200), "white").save(pdf, "PDF")
    message = TextDocConverter().validate(_job(pdf, "docx", tmp_path / "o.docx"))
    assert message is not None and "OCR" in message and "install" in message


def test_validate_scanned_pdf_passes_with_ocr_chain(tmp_path, monkeypatch):
    from PIL import Image

    _mock_ocr_chain(monkeypatch)
    pdf = tmp_path / "scan.pdf"
    Image.new("RGB", (200, 200), "white").save(pdf, "PDF")
    assert TextDocConverter().validate(_job(pdf, "docx", tmp_path / "o.docx")) is None


def test_validate_text_pdf_needs_no_ocr(tmp_path, monkeypatch):
    _mock_ocr_chain(monkeypatch, present=False)
    pdf = tmp_path / "text.pdf"
    _write_text_pdf(pdf, "Hello Converter, this is real text.")
    assert TextDocConverter().validate(_job(pdf, "docx", tmp_path / "o.docx")) is None


def test_broken_pdf_is_refused_cleanly(tmp_path):
    src = tmp_path / "bad.pdf"
    src.write_bytes(b"%PDF-1.4 not really a pdf")
    out = tmp_path / "out.docx"
    result = run_job(_job(src, "docx", out))
    assert result.ok is False
    assert "Could not read this PDF" in result.error
    assert not out.exists()


def test_ocr_gate_reads_limited_pages(monkeypatch):
    pytest.importorskip("pdfminer")
    import pdfminer.high_level as hl

    captured: dict = {}

    def fake_extract(path, **kw):
        captured.update(kw)
        return "questo pdf ha del testo reale " * 3

    monkeypatch.setattr(hl, "extract_text", fake_extract)
    assert _ocr.has_extractable_text(Path("x.pdf")) is True
    assert captured.get("maxpages") == _ocr.SCAN_CHECK_MAX_PAGES


# --- e2e: il PDF ricco ------------------------------------------------------

def test_pdf_to_markdown_keeps_headings_table_and_list(rich_pdf, tmp_path):
    out = tmp_path / "out.md"
    result = run_job(_job(rich_pdf, "md", out))
    assert result.ok, result.error
    markdown = out.read_text(encoding="utf-8")
    assert "# Quarterly report 2026" in markdown
    assert "# 1. Summary" in markdown
    assert "## 2. Figures" in markdown
    assert "| Region | Revenue | Growth |" in markdown
    assert "| North | 1,200 | 8% |" in markdown
    # La prima riga è l'header GFM, non una riga vuota aggiunta a tavolino.
    assert "|  |  |  |" not in markdown
    assert markdown.index("| Region | Revenue | Growth |") < markdown.index("| --- |")
    assert "- Expand the South region" in markdown
    assert "- Ship v2 in October" in markdown


def test_pdf_to_txt_keeps_reading_order_and_table_rows(rich_pdf, tmp_path):
    out = tmp_path / "out.txt"
    result = run_job(_job(rich_pdf, "txt", out))
    assert result.ok, result.error
    text = out.read_text(encoding="utf-8")
    assert "Region\tRevenue\tGrowth" in text
    assert "South\t940\t15%" in text
    # L'ordine di lettura: l'heading della tabella viene PRIMA delle celle
    # (pdfminer, da solo, le metteva prima l'heading).
    assert text.index("2. Figures") < text.index("Region")
    assert "- Hire two engineers" in text


def test_pdf_to_docx_builds_real_word_structure(rich_pdf, tmp_path):
    from docx import Document as Docx

    out = tmp_path / "out.docx"
    result = run_job(_job(rich_pdf, "docx", out, docx_layout="structure"))
    assert result.ok, result.error

    docx = Docx(str(out))
    styles = [(p.style.name, p.text) for p in docx.paragraphs if p.text.strip()]
    assert ("Title", REPORT) in styles
    assert ("Heading 1", "1. Summary") in styles
    assert ("Heading 2", "2. Figures") in styles
    assert any(style == "List Bullet" and text == "Hire two engineers"
               for style, text in styles)
    assert len(docx.tables) == 1
    grid = [[cell.text for cell in row.cells] for row in docx.tables[0].rows]
    assert grid == [["Region", "Revenue", "Growth"],
                    ["North", "1,200", "8%"],
                    ["South", "940", "15%"]]
    # La formattazione dell'originale viaggia con la struttura.
    heading = next(p for p in docx.paragraphs if p.text == "1. Summary")
    assert heading.runs[0].font.size.pt == pytest.approx(14.0)
    assert heading.runs[0].font.name == "Calibri"


def test_pdf_to_docx_visual_layout_uses_absolute_frames(rich_pdf, tmp_path):
    from docx import Document as Docx
    from docx.oxml.ns import qn

    out = tmp_path / "visual.docx"
    result = run_job(_job(rich_pdf, "docx", out, docx_layout="visual"))
    assert result.ok, result.error

    docx = Docx(str(out))
    framed = []
    for paragraph in docx.paragraphs:
        p_pr = paragraph._p.pPr
        frame = p_pr.find(qn("w:framePr")) if p_pr is not None else None
        if frame is not None and paragraph.text.strip():
            framed.append((paragraph.text.strip(), int(frame.get(qn("w:y")))))
    texts = [text for text, _ in framed]
    assert REPORT in texts
    assert any("Ship v2 in October" in text for text in texts)  # il proiettile torna nel layout visivo
    # Righe diverse stanno a quote diverse: non è un flusso, è un layout.
    assert len({y for _, y in framed}) > 5
    # Anche la pagina 2 è posizionata (e separata da un salto pagina).
    assert "Appendix" in texts
    xml = docx.element.xml
    assert xml.count("w:type=\"page\"") >= 1


def test_pdf_to_html_is_semantic(rich_pdf, tmp_path):
    out = tmp_path / "out.html"
    result = run_job(_job(rich_pdf, "html", out))
    assert result.ok, result.error
    page = out.read_text(encoding="utf-8")
    assert "<h1>" in page and "1. Summary" in page
    assert "<table" in page and "<li>" in page
    assert "position:absolute" not in page   # non è più il dump pdfminer


def test_pdf_to_odt_via_libreoffice(rich_pdf, tmp_path):
    if not deps.find_soffice():
        pytest.skip("LibreOffice required")
    out = tmp_path / "out.odt"
    result = run_job(_job(rich_pdf, "odt", out))
    assert result.ok, result.error
    assert b"mimetype" in out.read_bytes()[:200]


def test_docx_to_markdown_uses_gfm_tables(rich_pdf, tmp_path):
    """Il lettore DOCX + scrittore Markdown: tabelle pipe, non grid table."""
    from docx import Document as Docx

    src = tmp_path / "in.docx"
    doc = Docx()
    doc.add_heading("Title here", 0)
    doc.add_heading("Section", level=1)
    doc.add_paragraph("Body text with ")
    doc.paragraphs[-1].add_run("bold").bold = True
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(src)

    out = tmp_path / "out.md"
    result = run_job(_job(src, "md", out))
    assert result.ok, result.error
    markdown = out.read_text(encoding="utf-8")
    assert "# Title here" in markdown
    assert "# Section" in markdown
    assert "**bold**" in markdown
    assert "| A | B |" in markdown and "| 1 | 2 |" in markdown
    assert "-----" not in markdown   # niente grid table di pandoc


# --- Logica pura dell'IR (nessun tool esterno) ------------------------------

def test_split_font_extracts_family_and_style():
    assert _split_font("BAAAAA+Calibri-Bold") == ("Calibri", True, False)
    assert _split_font("DAAAAA+Cambria") == ("Cambria", False, False)
    assert _split_font("X+TimesNewRoman,BoldItalic") == ("Times New Roman", True, True)
    assert _split_font("FAAAAA+SymbolMT")[0] == ""


def test_xy_cut_orders_columns_left_then_right():
    def line(text, x, y):
        return Line(runs=[Run(text=text)], x=x, y=y, width=50, height=10)

    lines = [line("right top", 320, 700), line("left top", 50, 700),
             line("right bottom", 320, 680), line("left bottom", 50, 680),
             line("full width title", 50, 730)]
    ordered = [ln.text for ln in _xy_cut(lines)]
    assert ordered == ["full width title", "left top", "left bottom",
                       "right top", "right bottom"]


def test_xy_cut_splits_a_full_width_title_from_two_columns():
    """Un titolo che attraversa le colonne non deve mescolare le due colonne."""
    def line(text, x, y, width=50):
        return Line(runs=[Run(text=text)], x=x, y=y, width=width, height=10)

    lines = [line("right top", 320, 700), line("left top", 50, 700),
             line("right bottom", 320, 680), line("left bottom", 50, 680),
             line("full width title", 50, 730, width=500)]
    ordered = [ln.text for ln in _xy_cut(lines)]
    assert ordered == ["full width title", "left top", "left bottom",
                       "right top", "right bottom"]


def test_paragraph_text_dehyphenates_and_rejoins_lines():
    paragraph = Paragraph(lines=[
        Line(runs=[Run(text="inter-")], x=0, y=100, width=40, height=10),
        Line(runs=[Run(text="national standards")], x=0, y=88, width=120, height=10),
    ])
    assert paragraph.text == "international standards"


def test_document_body_size_is_weighted_by_characters():
    document = Document(blocks=[
        Paragraph(lines=[Line(runs=[Run(text="x" * 500, size=11.0)])]),
        Paragraph(lines=[Line(runs=[Run(text="heading", size=20.0)])]),
    ])
    assert document.body_size() == 11.0


# --- Robustezza dei lettori (casi reali del corpus) -------------------------

def _write_positioned_pdf(path: Path, content: bytes) -> None:
    """PDF con due parole posizionate: il content stream non contiene lo spazio."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode()
            + b" /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF\n")
    path.write_bytes(out)


def test_pdf_keeps_spaces_between_positioned_words(tmp_path):
    """pdfminer mette gli spazi sintetici in LTAnno, non in LTChar."""
    from converter.converters.docmodel import read_pdf

    pdf = tmp_path / "spaced.pdf"
    _write_positioned_pdf(
        pdf, b"BT /F1 18 Tf 72 720 Td (Hello) Tj 60 0 Td (world) Tj ET")
    document = read_pdf(pdf)
    assert [block.text for block in document.blocks] == ["Hello world"]


def _write_placed_text_pdf(path: Path, items: list[tuple[float, float, str]],
                           *, rotated: list[tuple[float, float, str]] = (),
                           rotated_size: float = 14.0) -> None:
    """PDF minimo con testo posizionato: (x, y, testo) in punti tipografici.

    `rotated` scrive il testo ruotato di 90° (un glifo ogni 10pt verso il basso),
    come i timbri "arXiv:..." nel margine sinistro dei PDF arXiv.
    """
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    parts = ["BT /F1 10 Tf\n"]
    parts += [f"1 0 0 1 {x:.1f} {y:.1f} Tm ({escape(text)}) Tj\n" for x, y, text in items]
    parts.append("ET")
    if rotated:
        parts.append(f"\nBT /F1 {rotated_size:.1f} Tf\n")
        parts += [f"0 1 -1 0 {x:.1f} {y - 10.0 * index:.1f} Tm ({escape(char)}) Tj\n"
                  for x, y, text in rotated for index, char in enumerate(text)]
        parts.append("ET")
    _write_positioned_pdf(path, "".join(parts).encode("latin-1"))


def test_pdf_aligned_numeric_table_is_detected(tmp_path):
    """Una tabella senza linee: le colonne sono solo parole allineate."""
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Table

    rows = [("200.01", "1"), ("400.01", "3"), ("600.01", "5"), ("800.01", "7")]
    items: list[tuple[float, float, str]] = []
    for index, (amount, count) in enumerate(rows):
        items.append((72.0, 700.0 - 16 * index, amount))
        items.append((240.0, 700.0 - 16 * index, count))
    pdf = tmp_path / "numbers.pdf"
    _write_placed_text_pdf(pdf, items)

    document = read_pdf(pdf)
    tables = [block for block in document.blocks if isinstance(block, Table)]
    assert len(tables) == 1
    assert tables[0].borders is False
    assert [[cell.paragraphs[0].text for cell in row] for row in tables[0].rows] == [
        ["200.01", "1"], ["400.01", "3"], ["600.01", "5"], ["800.01", "7"]]


def test_pdf_aligned_table_needs_a_caption_or_numbers(tmp_path):
    """Tre colonne di testo senza numeri sono una tabella solo con la didascalia."""
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Table

    rows = [("Alpha", "Beta", "Gamma"), ("One", "Two", "Three")]
    body: list[tuple[float, float, str]] = []
    for index, row in enumerate(rows):
        for column, text in enumerate(row):
            body.append((72.0 + 120 * column, 660.0 - 16 * index, text))

    without = tmp_path / "no_caption.pdf"
    _write_placed_text_pdf(without, body)
    assert not [block for block in read_pdf(without).blocks if isinstance(block, Table)]

    with_caption = tmp_path / "caption.pdf"
    _write_placed_text_pdf(with_caption, [(72.0, 700.0, "Table 1: Letters")] + body)
    tables = [block for block in read_pdf(with_caption).blocks if isinstance(block, Table)]
    assert len(tables) == 1
    assert [[cell.paragraphs[0].text for cell in row] for row in tables[0].rows] == [
        ["Alpha", "Beta", "Gamma"], ["One", "Two", "Three"]]


def test_pdf_two_columns_of_text_are_not_a_table(tmp_path):
    """Due colonne di prosa hanno celle lunghe e nessun numero: non è una tabella."""
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Table

    items: list[tuple[float, float, str]] = []
    for index in range(4):
        items.append((72.0, 700.0 - 14 * index, "This is a long sentence in the left column"))
        items.append((330.0, 700.0 - 14 * index, "And this is the right column text here"))
    pdf = tmp_path / "columns.pdf"
    _write_placed_text_pdf(pdf, items)

    document = read_pdf(pdf)
    assert not [block for block in document.blocks if isinstance(block, Table)]
    assert any("left column" in block.text for block in document.blocks)


def test_pdf_table_beside_a_text_column_keeps_prose_out_of_the_table(tmp_path):
    """La prosa della colonna sinistra non deve diventare una colonna della tabella."""
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Table

    prose = ["Left column prose line number one",
             "Left column prose line number two",
             "Left column prose line number three",
             "Left column prose line number four"]
    items: list[tuple[float, float, str]] = []
    for index, text in enumerate(prose):
        items.append((72.0, 700.0 - 16 * index, text))
        items.append((330.0, 700.0 - 16 * index, str(index + 1)))
        items.append((420.0, 700.0 - 16 * index, str(10 * (index + 1))))
    pdf = tmp_path / "table_beside_text.pdf"
    _write_placed_text_pdf(pdf, items)

    document = read_pdf(pdf)
    tables = [block for block in document.blocks if isinstance(block, Table)]
    assert len(tables) == 1
    assert [[cell.paragraphs[0].text for cell in row] for row in tables[0].rows] == [
        ["1", "10"], ["2", "20"], ["3", "30"], ["4", "40"]]
    body = [block.text for block in document.blocks if not isinstance(block, Table)]
    assert any("Left column prose line number one" in text for text in body)


def test_pdf_wide_gap_before_a_description_column_keeps_the_column(tmp_path):
    """Una colonna di descrizioni è una colonna della tabella, non prosa da scartare."""
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Table

    rows = [("Line 1", "01-05", "Command Code required for the display"),
            ("Line 2", "06-15", "Command Definer Code is required"),
            ("Line 3", "17-20", "Display Source Code is required"),
            ("Line 4", "21-30", "Name Control of the taxpayer")]
    items: list[tuple[float, float, str]] = [(72.0, 730.0, "Table 2: Command Code Input")]
    for index, row in enumerate(rows):
        items.append((72.0, 700.0 - 16 * index, row[0]))
        items.append((110.0, 700.0 - 16 * index, row[1]))
        items.append((300.0, 700.0 - 16 * index, row[2]))
    pdf = tmp_path / "wide_gap_table.pdf"
    _write_placed_text_pdf(pdf, items)

    document = read_pdf(pdf)
    tables = [block for block in document.blocks if isinstance(block, Table)]
    assert len(tables) == 1
    assert [[cell.paragraphs[0].text for cell in row] for row in tables[0].rows] == [
        list(row) for row in rows]


def test_pdf_merge_fragments_does_not_join_columns_or_mutate_input():
    """Un vuoto di 12pt tra due colonne non è uno spazio: le righe restano separate."""
    from converter.converters.docmodel.pdfreader import _merge_fragments

    left = Line(runs=[Run(text="left column")], x=72, y=700, width=60, height=10)
    right = Line(runs=[Run(text="right column")], x=144, y=700, width=65, height=10)
    merged = _merge_fragments([left, right])
    assert [line.text for line in merged] == ["left column", "right column"]
    assert left.text == "left column"       # l'ingresso non viene mutato

    bullet = Line(runs=[Run(text="●")], x=72, y=680, width=5, height=10)
    text = Line(runs=[Run(text="item text")], x=90, y=680, width=40, height=10)
    assert [line.text for line in _merge_fragments([bullet, text])] == ["● item text"]


def test_pdf_rotated_margin_stamp_stays_out_of_the_body(tmp_path):
    """Il timbro arXiv ruotato non si fonde col corpo né diventa un titolo.

    pdfminer spezza il testo ruotato in un glifo per riga: la cifra sulla stessa
    base veniva assorbita come proiettile e i glifi più grandi superavano la
    soglia heading.
    """
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Paragraph

    pdf = tmp_path / "rotated_stamp.pdf"
    _write_placed_text_pdf(
        pdf,
        [(72.0, 700.0, "Body text of the document."),
         (72.0, 684.0, "A second line of the same paragraph.")],
        rotated=[(16.0, 700.0, "4arXiv:2004.04906v3")],
        rotated_size=20.0,
    )

    document = read_pdf(pdf)
    paragraphs = [block for block in document.blocks if isinstance(block, Paragraph)]
    texts = [paragraph.text for paragraph in paragraphs]
    assert any("Body text of the document." in text for text in texts)
    assert all("arXiv" not in text for text in texts)
    assert all(not text.startswith("4") for text in texts)
    assert all(paragraph.level == 0 for paragraph in paragraphs)


def test_pdf_numbered_heading_slightly_larger_than_body_stays_separate(tmp_path):
    """Un titolo numerato non bold appena più grande non si fonde col paragrafo.

    "1 Introduction" a 11pt su un corpo di 10pt non supera la soglia 1.15× del
    rilevamento heading e la differenza di 1pt non basta a fermare il merge:
    senza la regola dedicata la riga diventava l'inizio del paragrafo.
    """
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.types import Paragraph

    pdf = tmp_path / "numbered_heading.pdf"
    body = ("Open-domain question answering answers factoid questions using a large "
            "collection of documents and a retriever.")
    content = (
        "BT /F1 11 Tf 1 0 0 1 72 700 Tm (1 Introduction) Tj ET\n"
        f"BT /F1 10 Tf 1 0 0 1 72 684 Tm ({body}) Tj ET\n"
        "BT /F1 10 Tf 1 0 0 1 72 670 Tm (The second line belongs to the same paragraph.) Tj ET"
    )
    _write_positioned_pdf(pdf, content.encode("latin-1"))

    document = read_pdf(pdf)
    paragraphs = [block for block in document.blocks if isinstance(block, Paragraph)]
    texts = [paragraph.text for paragraph in paragraphs]
    assert "1 Introduction" in texts
    assert any(text.startswith("Open-domain question answering") for text in texts)
    heading = next(paragraph for paragraph in paragraphs if paragraph.text == "1 Introduction")
    assert heading.style == "heading"
    assert heading.level == 1


def test_pdf_images_are_stored_as_compact_png():
    """pdfminer esporta BMP grezzi: l'IR li converte in PNG per non gonfiare i writer."""
    import io

    from PIL import Image

    from converter.converters.docmodel.pdfreader import _compact_image

    buffer = io.BytesIO()
    Image.new("RGB", (120, 80), (40, 90, 200)).save(buffer, format="BMP")
    bmp = buffer.getvalue()
    data, ext = _compact_image(bmp, "bmp")
    assert ext == "png"
    assert data.startswith(b"\x89PNG\r\n\x1a\n")
    assert len(data) < len(bmp)

    # Un formato già compresso non viene ricodificato.
    png_buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (0, 0, 0)).save(png_buffer, format="PNG")
    png = png_buffer.getvalue()
    assert _compact_image(png, "png") == (png, "png")

    # Flusso illeggibile: si tiene l'originale invece di perdere l'immagine.
    assert _compact_image(b"not an image", "bmp") == (b"not an image", "bmp")


def _set_numbering(paragraph, num_id: int, level: int = 0) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def test_docx_lists_come_from_numpr_not_style_name(tmp_path):
    """Le liste Word reali sono `List Paragraph` + w:numPr, non `List Bullet`."""
    from docx import Document as Docx

    from converter.converters.docmodel import read_docx
    from converter.converters.docmodel.types import ListBlock

    src = tmp_path / "lists.docx"
    doc = Docx()
    for text, num_id, level in (("bullet one", 1, 0), ("bullet two", 1, 0),
                                ("nested", 1, 1), ("number one", 5, 0)):
        paragraph = doc.add_paragraph(text, style="List Paragraph")
        _set_numbering(paragraph, num_id, level)
    doc.save(src)

    document = read_docx(src)
    lists = [block for block in document.blocks if isinstance(block, ListBlock)]
    assert [(block.ordered, [(item.level, item.paragraph.text) for item in block.items])
            for block in lists] == [
        (False, [(1, "bullet one"), (1, "bullet two"), (2, "nested")]),
        (True, [(1, "number one")]),
    ]


def test_docx_ignores_comments_in_body(tmp_path):
    """Un commento nel body non è un paragrafo: lxml gli dà una funzione come tag."""
    from docx import Document as Docx
    from lxml import etree

    from converter.converters.docmodel import read_docx

    src = tmp_path / "comment.docx"
    doc = Docx()
    doc.add_paragraph("Before")
    doc.element.body.append(etree.Comment("not a paragraph"))
    doc.add_paragraph("After")
    doc.save(src)

    document = read_docx(src)
    assert [block.text for block in document.blocks] == ["Before", "After"]


def test_docx_merged_cells_are_not_duplicated(tmp_path):
    """Una cella unita su più colonne è una cella, non N copie con lo stesso testo."""
    from docx import Document as Docx

    from converter.converters.docmodel import read_docx
    from converter.converters.docmodel.types import Table

    src = tmp_path / "merged.docx"
    doc = Docx()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Merged header"
    for index, value in enumerate(("A", "B", "C")):
        table.cell(1, index).text = value
    doc.save(src)

    document = read_docx(src)
    block = next(b for b in document.blocks if isinstance(b, Table))
    assert [[cell.paragraphs[0].text for cell in row] for row in block.rows] == [
        ["Merged header"], ["A", "B", "C"],
    ]


def test_docx_without_page_size_or_table_grid_does_not_crash(tmp_path):
    from docx import Document as Docx
    from docx.oxml.ns import qn

    from converter.converters.docmodel import read_docx

    src = tmp_path / "broken.docx"
    doc = Docx()
    doc.add_paragraph("Text survives")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"

    sect_pr = doc.sections[0]._sectPr
    for tag in ("w:pgSz", "w:pgMar"):
        element = sect_pr.find(qn(tag))
        if element is not None:
            sect_pr.remove(element)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        table._tbl.remove(grid)
    doc.save(src)

    document = read_docx(src)
    assert document.blocks[0].text == "Text survives"
    assert document.blocks[1].rows[0][0].paragraphs[0].text == "A"
