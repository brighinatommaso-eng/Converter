"""Test end-to-end del converter dati (un converter per categoria)."""

import json

import pandas as pd
import pytest
import yaml

import converter.converters  # noqa: F401
from converter.core.job import ConversionJob
from converter.core.pipeline import run_job
from converter.core.registry import resolve

# Classico "billion laughs": DTD con entità che si espandono esponenzialmente.
_XML_BOMB = """<?xml version="1.0"?>
<!DOCTYPE lolz [
 <!ENTITY lol "lol">
 <!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">
 <!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">
]>
<lolz>&lol3;</lolz>"""


def test_xml_to_csv_works(tmp_path):
    src = tmp_path / "d.xml"
    out = tmp_path / "d.csv"
    src.write_text("<data><row><a>1</a><b>2</b></row><row><a>3</a><b>4</b></row></data>")
    result = resolve("xml", "csv").convert(ConversionJob(src, "xml", "csv", out), lambda e: None)
    assert result.ok and out.exists()
    assert "a,b" in out.read_text()


def test_xml_bomb_is_rejected_gracefully(tmp_path):
    src = tmp_path / "bomb.xml"
    out = tmp_path / "bomb.csv"
    src.write_text(_XML_BOMB)
    # Deve fallire in modo pulito (niente crash, niente esplosione di memoria).
    result = run_job(ConversionJob(src, "xml", "csv", out))
    assert not result.ok
    assert "xml" in result.error.lower()
    assert not out.exists()


def test_csv_to_json_roundtrips_values(tmp_path):
    src = tmp_path / "d.csv"
    out = tmp_path / "d.json"
    src.write_text("name,age\nAlice,30\nBob,25\n")
    conv = resolve("csv", "json")
    result = conv.convert(
        ConversionJob(src, "csv", "json", out, options={"pretty": True, "orient": "records"}),
        lambda e: None,
    )
    assert result.ok and out.exists()
    data = json.loads(out.read_text())
    assert data[0]["name"] == "Alice"
    assert int(data[0]["age"]) == 30
    assert len(data) == 2


def test_csv_to_tsv_is_lossless(tmp_path):
    src = tmp_path / "d.csv"
    out = tmp_path / "d.tsv"
    src.write_text("a,b\n1,2\n3,4\n")
    conv = resolve("csv", "tsv")
    result = conv.convert(
        ConversionJob(src, "csv", "tsv", out, options={"header": True}), lambda e: None
    )
    assert result.ok and out.exists()
    assert result.verdict.is_lossless
    content = out.read_text()
    assert "\t" in content and "a\tb" in content


# --- Bug 1: json<->yaml e json->json dichiarati LOSSLESS ma pandas appiattiva ---

def test_json_to_yaml_preserves_nesting(tmp_path):
    """Regressione BUG 1: json->yaml passava da json_normalize e distruggeva
    l'annidamento (chiavi appiattite col punto, oggetto avvolto in lista)."""
    src = tmp_path / "d.json"
    out = tmp_path / "d.yaml"
    payload = {"user": {"name": "Ada", "roles": ["a", "b"]}, "n": 1}
    src.write_text(json.dumps(payload))
    conv = resolve("json", "yaml")
    result = conv.convert(ConversionJob(src, "json", "yaml", out), lambda e: None)
    assert result.ok and out.exists()
    assert result.verdict.is_lossless
    assert yaml.safe_load(out.read_text()) == payload


def test_json_to_json_preserves_object(tmp_path):
    """Regressione BUG 1: json->json ('Same format.' LOSSLESS) trasformava
    l'oggetto in una LISTA di record appiattiti."""
    src = tmp_path / "d.json"
    out = tmp_path / "o.json"
    payload = {"user": {"name": "Ada", "roles": ["a", "b"]}, "n": 1}
    src.write_text(json.dumps(payload))
    result = resolve("json", "json").convert(
        ConversionJob(src, "json", "json", out, options={"pretty": True}), lambda e: None
    )
    assert result.ok and out.exists()
    assert json.loads(out.read_text()) == payload


def test_yaml_to_json_preserves_nesting(tmp_path):
    src = tmp_path / "d.yaml"
    out = tmp_path / "d.json"
    payload = {"user": {"name": "Ada", "roles": ["a", "b"]}, "n": 1}
    src.write_text(yaml.safe_dump(payload))
    result = resolve("yaml", "json").convert(
        ConversionJob(src, "yaml", "json", out), lambda e: None
    )
    assert result.ok and out.exists()
    assert json.loads(out.read_text()) == payload


# --- Bug 2: csv<->tsv dichiarato LOSSLESS ma pandas re-inferiva i tipi ---

def test_csv_to_tsv_preserves_leading_zeros_and_decimals(tmp_path):
    """Regressione BUG 2: pd.read_csv senza dtype inferiva i tipi
    (01234->1234, 1.10->1.1, 00500->500)."""
    src = tmp_path / "d.csv"
    out = tmp_path / "d.tsv"
    src.write_text("code,price\n01234,1.10\n00500,2.00\n")
    result = resolve("csv", "tsv").convert(
        ConversionJob(src, "csv", "tsv", out, options={"header": True}), lambda e: None
    )
    assert result.ok and out.exists()
    assert result.verdict.is_lossless
    lines = out.read_text().splitlines()
    assert lines[0] == "code\tprice"
    assert lines[1] == "01234\t1.10"
    assert lines[2] == "00500\t2.00"


def test_csv_to_xlsx_preserves_values(tmp_path):
    """Regressione BUG 2 (parte xlsx/ods): il verdetto csv→xlsx è LOSSLESS
    'maps faithfully' — deve esserlo davvero, niente inferenza dei tipi."""
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    src = tmp_path / "d.csv"
    out = tmp_path / "d.xlsx"
    src.write_text("code,price\n01234,1.10\n00500,2.00\n")
    result = resolve("csv", "xlsx").convert(ConversionJob(src, "csv", "xlsx", out), lambda e: None)
    assert result.ok and out.exists()
    assert result.verdict.is_lossless
    ws = load_workbook(out).active
    assert ws["A2"].value == "01234"  # non 1234
    assert ws["B2"].value == "1.10"   # non 1.1
    assert ws["A3"].value == "00500"


def test_csv_to_json_preserves_values(tmp_path):
    """csv→json è LOSSLESS 'maps faithfully': i valori testuali del CSV vanno
    preservati esattamente (un CSV non ha tipi)."""
    src = tmp_path / "d.csv"
    out = tmp_path / "d.json"
    src.write_text("code,price\n01234,1.10\n")
    result = resolve("csv", "json").convert(
        ConversionJob(src, "csv", "json", out, options={"pretty": False, "orient": "records"}),
        lambda e: None,
    )
    assert result.ok and out.exists()
    data = json.loads(out.read_text())
    assert data[0]["code"] == "01234" and data[0]["price"] == "1.10"


def test_same_format_xlsx_preserves_sheets_and_formulas(tmp_path):
    """Regressione: xlsx→xlsx ha verdict LOSSLESS 'Same format.' — deve esserlo
    davvero. Il vecchio round-trip pandas teneva solo i valori del PRIMO foglio
    (formule e fogli extra persi) pur dichiarandosi lossless. Ora è una copia
    fedele: tutti i fogli e le formule sopravvivono."""
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    src = tmp_path / "book.xlsx"
    out = tmp_path / "out.xlsx"
    with pd.ExcelWriter(src, engine="openpyxl") as w:
        pd.DataFrame({"a": [1, 2], "b": [3, 4]}).to_excel(w, sheet_name="Sheet1", index=False)
        pd.DataFrame({"z": [9]}).to_excel(w, sheet_name="Extra", index=False)
    wb = load_workbook(src)
    wb["Sheet1"]["D1"] = "=A2+B2"
    wb.save(src)

    result = resolve("xlsx", "xlsx").convert(ConversionJob(src, "xlsx", "xlsx", out), lambda e: None)
    assert result.ok and result.verdict.is_lossless
    got = load_workbook(out)
    assert got.sheetnames == ["Sheet1", "Extra"]      # niente foglio perso
    assert got["Sheet1"]["D1"].value == "=A2+B2"      # formula preservata


def test_same_format_xml_is_byte_identical(tmp_path):
    """xml→xml LOSSLESS: copia fedele, non un round-trip pandas approssimato."""
    src = tmp_path / "d.xml"
    out = tmp_path / "out.xml"
    src.write_bytes(b'<?xml version="1.0"?>\n<root attr="1"><item>x</item></root>\n')
    result = resolve("xml", "xml").convert(ConversionJob(src, "xml", "xml", out), lambda e: None)
    assert result.ok and result.verdict.is_lossless
    assert out.read_bytes() == src.read_bytes()


def test_same_format_xlsx_strips_metadata_only_on_request(tmp_path):
    """Con byte-copy i metadati del documento sopravvivono (l'utente non ha
    chiesto lo strip) — coerente con le immagini; con strip_metadata=True le
    proprietà vengono rimosse ma i dati restano."""
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook, load_workbook

    src = tmp_path / "doc.xlsx"
    wb = Workbook()
    wb.properties.creator = "SECRET-AUTHOR"
    wb.properties.title = "SECRET-DOC-TITLE"
    wb.active.append(["name", "age"])
    wb.active.append(["Alice", 30])
    wb.save(src)

    # Strip OFF: i metadati sono preservati (copia fedele).
    keep = tmp_path / "keep.xlsx"
    resolve("xlsx", "xlsx").convert(ConversionJob(src, "xlsx", "xlsx", keep), lambda e: None)
    assert load_workbook(keep).properties.creator == "SECRET-AUTHOR"

    # Strip ON: le proprietà vengono rimosse, i dati restano.
    strip = tmp_path / "strip.xlsx"
    resolve("xlsx", "xlsx").convert(
        ConversionJob(src, "xlsx", "xlsx", strip, strip_metadata=True), lambda e: None)
    props = load_workbook(strip).properties
    assert props.creator != "SECRET-AUTHOR"
    assert props.title != "SECRET-DOC-TITLE"
    assert [c.value for c in load_workbook(strip).active[2]] == ["Alice", 30]  # dati intatti


# --- Badge LOSSLESS falsi (review 2026-10-26) --------------------------------

def _job(src, s, d, out, **opts):
    return ConversionJob(src_path=src, src_format=s, dst_format=d, output_path=out, options=opts)


def test_yaml_to_yaml_is_a_byte_copy(tmp_path):
    """Il re-dump perdeva i commenti e cambiava `NO` in `false` sotto LOSSLESS."""
    src = tmp_path / "d.yaml"
    src.write_text("# commento importante\ncountry: NO\nreleased: 2024-01-01\n", encoding="utf-8")
    out = tmp_path / "o.yaml"
    result = run_job(_job(src, "yaml", "yaml", out))
    assert result.ok and result.verdict.is_lossless
    assert out.read_bytes() == src.read_bytes()


def test_yaml_to_json_handles_dates_and_is_honestly_lossy(tmp_path):
    src = tmp_path / "d.yaml"
    src.write_text("# c\nreleased: 2024-01-01\n", encoding="utf-8")
    out = tmp_path / "d.json"
    result = run_job(_job(src, "yaml", "json", out))
    assert result.ok, result.error
    assert result.verdict.is_lossy and "comments" in result.verdict.reason
    assert json.loads(out.read_text(encoding="utf-8")) == {"released": "2024-01-01"}


def test_semicolon_csv_keeps_its_columns(tmp_path):
    """I CSV europei usano ';': letti con ',' diventavano UNA colonna (LOSSLESS)."""
    src = tmp_path / "prezzi.csv"
    src.write_text("nome;prezzo\nmela;1,50\npera;2,00\n", encoding="utf-8")
    out = tmp_path / "prezzi.xlsx"
    assert run_job(_job(src, "csv", "xlsx", out)).ok
    df = pd.read_excel(out, dtype=str)
    assert list(df.columns) == ["nome", "prezzo"]
    assert df["prezzo"].tolist() == ["1,50", "2,00"]
    # csv→csv conserva il separatore della sorgente.
    same = tmp_path / "copia.csv"
    assert run_job(_job(src, "csv", "csv", same)).ok
    assert same.read_text(encoding="utf-8").splitlines()[0] == "nome;prezzo"


def test_cp1252_csv_is_read_instead_of_crashing(tmp_path):
    src = tmp_path / "excel.csv"
    src.write_bytes("città,prezzo\nForlì,3\n".encode("cp1252"))
    out = tmp_path / "excel.json"
    result = run_job(_job(src, "csv", "json", out))
    assert result.ok, result.error
    assert json.loads(out.read_text(encoding="utf-8")) == [{"città": "Forlì", "prezzo": "3"}]


def test_duplicate_headers_survive_tabular_targets(tmp_path):
    """pandas rinominava "a,a" in "a,a.1" sotto un badge LOSSLESS."""
    src = tmp_path / "dup.csv"
    src.write_text("a,a,b\n1,2,3\n", encoding="utf-8")
    out = tmp_path / "dup.tsv"
    result = run_job(_job(src, "csv", "tsv", out))
    assert result.ok and result.verdict.is_lossless
    assert out.read_text(encoding="utf-8").splitlines() == ["a\ta\tb", "1\t2\t3"]


def test_duplicate_headers_to_json_are_renamed_and_declared(tmp_path):
    src = tmp_path / "dup.csv"
    src.write_text("a,a,b\n1,2,3\n", encoding="utf-8")
    out = tmp_path / "dup.json"
    result = run_job(_job(src, "csv", "json", out))
    assert result.ok, result.error
    assert result.verdict.is_lossy and "'a' → 'a_2'" in result.verdict.reason
    assert json.loads(out.read_text(encoding="utf-8")) == [{"a": "1", "a_2": "2", "b": "3"}]


def test_xlsx_to_json_still_writes_numbers(tmp_path):
    src = tmp_path / "n.xlsx"
    pd.DataFrame({"n": [1, 2], "x": [1.5, None]}).to_excel(src, index=False)
    out = tmp_path / "n.json"
    result = run_job(_job(src, "xlsx", "json", out))
    assert result.ok, result.error
    assert json.loads(out.read_text(encoding="utf-8")) == [{"n": 1, "x": 1.5}, {"n": 2, "x": None}]


def test_duplicate_headers_to_xml_stay_valid_tag_names(tmp_path):
    src = tmp_path / "dup.csv"
    src.write_text("a,a\n1,2\n", encoding="utf-8")
    out = tmp_path / "dup.xml"
    result = run_job(_job(src, "csv", "xml", out))
    assert result.ok, result.error
    assert "<a_2>2</a_2>" in out.read_text(encoding="utf-8")
