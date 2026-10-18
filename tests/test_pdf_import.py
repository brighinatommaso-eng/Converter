"""PdfImportConverter: routing, disponibilità per modalità/dipendenza, e il
rifiuto onesto dei PDF scansionati in validate() (mai un docx vuoto "riuscito")."""

from pathlib import Path

import pytest
from PIL import Image

from converter.converters import pdf_import
from converter.converters.pdf_import import PdfImportConverter
from converter.core import deps, engine, registry
from converter.core.job import ConversionJob


def _mock_ocr_chain(monkeypatch, *, ocrmypdf=True, ghostscript=True, tesseract=True):
    """Finge presente/assente la catena OCR (mocka le dipendenze, non la modalità)."""
    monkeypatch.setattr(deps, "find_ocrmypdf",
                        lambda: r"C:\Scripts\ocrmypdf.exe" if ocrmypdf else None)
    monkeypatch.setattr(deps, "find_ghostscript",
                        lambda: r"C:\gs\bin\gswin64c.exe" if ghostscript else None)
    monkeypatch.setattr(deps, "find_tesseract",
                        lambda: r"C:\Tesseract-OCR\tesseract.exe" if tesseract else None)


def _write_text_pdf(path: Path, text: str) -> None:
    """PDF minimale con testo VERO (operatore Tj), scritto a mano: Pillow sa
    scrivere solo PDF-immagine e qui serve l'opposto. xref calcolato."""
    content = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode()
            + b" /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF\n")
    path.write_bytes(out)


def _job(src: Path, dst_format: str, out: Path) -> ConversionJob:
    return ConversionJob(src_path=src, src_format="pdf",
                         dst_format=dst_format, output_path=out)


# --- Routing ---

def test_resolve_pdf_writer_targets_full():
    engine.set_mode("FULL")
    for dst in ("docx", "odt", "rtf"):
        conv = registry.resolve("pdf", dst)
        assert isinstance(conv, PdfImportConverter), (dst, type(conv))


# --- Disponibilità per modalità ---

def test_portable_pdf_writer_targets_greyed(monkeypatch):
    monkeypatch.setattr(deps, "find_soffice", lambda: r"C:\soffice.exe")
    engine.set_mode("PORTABLE")
    targets = {t.dst_format: t for t in registry.targets_for("pdf")}
    for dst in ("docx", "odt", "rtf"):
        assert targets[dst].available is False, dst
        assert "Full mode" in targets[dst].reason_unavailable, dst
        # La ragione è quella del NUOVO converter, non il pair-block di
        # DocumentConverter ("no export filter exists...").
        assert "export filter" not in targets[dst].reason_unavailable, dst


def test_portable_reason_is_not_install_hint():
    engine.set_mode("PORTABLE")
    ds = PdfImportConverter.available()
    assert ds.ok is False
    assert ds.install_hint == ""


# --- Disponibilità per dipendenza (mocka la dipendenza, non la modalità) ---

def test_full_without_pdf2docx_unavailable(monkeypatch):
    engine.set_mode("FULL")
    real = deps.has_module
    monkeypatch.setattr(deps, "has_module",
                        lambda name: False if name == "pdf2docx" else real(name))
    monkeypatch.setattr(deps, "find_soffice", lambda: r"C:\soffice.exe")
    ds = PdfImportConverter.available()
    assert ds.ok is False
    assert "pdf2docx" in ds.reason + ds.install_hint


def test_targets_pdf_writer_greyed_without_pdf2docx(monkeypatch):
    """Scenario release (PyMuPDF/pdf2docx NON nel bundle): pdf→docx/odt/rtf
    compaiono ma non disponibili, con un motivo chiaro e senza crash; i target
    PDF puri (txt/html, pdfminer) restano disponibili."""
    engine.set_mode("FULL")
    real = deps.has_module
    monkeypatch.setattr(deps, "has_module",
                        lambda name: False if name == "pdf2docx" else real(name))
    monkeypatch.setattr(deps, "find_soffice", lambda: r"C:\soffice.exe")
    targets = {t.dst_format: t for t in registry.targets_for("pdf")}
    for dst in ("docx", "odt", "rtf"):
        assert targets[dst].available is False, dst
        assert "pdf2docx" in targets[dst].reason_unavailable, dst
    assert targets["txt"].available is True
    assert targets["html"].available is True


def test_full_without_soffice_docx_available_odt_rtf_greyed(monkeypatch):
    # Risk 17: senza LibreOffice, pdf→docx (solo pdf2docx) resta disponibile;
    # solo pdf→odt/rtf (secondo salto LibreOffice) va in grigio, con il motivo giusto.
    engine.set_mode("FULL")
    monkeypatch.setattr(deps, "has_module", lambda name: True)
    monkeypatch.setattr(deps, "find_soffice", lambda: None)
    conv = PdfImportConverter()
    assert conv.available().ok is True
    assert conv.pair_available("pdf", "docx") == (True, "")
    for dst in ("odt", "rtf"):
        ok, reason = conv.pair_available("pdf", dst)
        assert ok is False and "LibreOffice" in reason


# --- Rilevamento scansione in validate() (ramo OCR degradante) ---

@pytest.mark.parametrize("missing", ["ocrmypdf", "ghostscript", "tesseract"])
def test_validate_scanned_pdf_refused_without_ocr(tmp_path, monkeypatch, missing):
    """Senza un pezzo QUALSIASI della catena OCR, la scansione è rifiutata col
    messaggio degradante che spiega cosa installare."""
    _mock_ocr_chain(monkeypatch, **{missing: False})
    pdf = tmp_path / "scan.pdf"
    Image.new("RGB", (300, 300), "white").save(pdf, "PDF")
    msg = PdfImportConverter().validate(_job(pdf, "docx", tmp_path / "o.docx"))
    assert msg is not None
    assert "scanned" in msg and "OCR" in msg and "install" in msg


def test_validate_scanned_pdf_passes_with_ocr(tmp_path, monkeypatch):
    """Con la catena OCR completa la scansione passa: l'OCR avviene in convert()."""
    _mock_ocr_chain(monkeypatch)
    pdf = tmp_path / "scan.pdf"
    Image.new("RGB", (300, 300), "white").save(pdf, "PDF")
    msg = PdfImportConverter().validate(_job(pdf, "docx", tmp_path / "o.docx"))
    assert msg is None


def test_validate_text_pdf_passes_without_ocr(tmp_path, monkeypatch):
    """DEGRADANTE: un PDF con testo si converte anche senza catena OCR."""
    _mock_ocr_chain(monkeypatch, ocrmypdf=False, ghostscript=False, tesseract=False)
    pdf = tmp_path / "text.pdf"
    _write_text_pdf(pdf, "Hello Converter, this is real extractable text.")
    msg = PdfImportConverter().validate(_job(pdf, "docx", tmp_path / "o.docx"))
    assert msg is None


def test_validate_text_pdf_passes(tmp_path):
    pdf = tmp_path / "text.pdf"
    _write_text_pdf(pdf, "Hello Converter, this is real extractable text.")
    msg = PdfImportConverter().validate(_job(pdf, "docx", tmp_path / "o.docx"))
    assert msg is None


# --- Ricostruzione docx dal testo OCR ---

def test_ocr_pdf_to_docx_paragraph_logic(tmp_path):
    """SOLA LOGICA del builder (paragrafi/pagine) su un PDF con testo visibile:
    NON esercita il layer invisibile — quello è coperto dal test successivo
    (con catena OCR reale) e dall'e2e."""
    from docx import Document

    from converter.converters.pdf_import import _ocr_pdf_to_docx

    pdf = tmp_path / "text.pdf"
    _write_text_pdf(pdf, "Hello Converter, this is real extractable text.")
    out = tmp_path / "o.docx"
    ok, err = _ocr_pdf_to_docx(pdf, out)
    assert ok is True, err
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Hello Converter" in text


# --- Rischio 9c: pdf2docx gira in un processo figlio killabile ------------------

def test_pdf_to_docx_subprocess_produces_docx(tmp_path):
    """Il salto pdf→docx passa da _pdf_to_docx, che lancia pdf2docx in un processo
    figlio (timeout duro + kill dell'albero) invece che in-process. Su un PDF
    testuale valido deve comunque produrre un docx non vuoto."""
    pytest.importorskip("pdf2docx")
    src = tmp_path / "t.pdf"
    _write_text_pdf(src, "Hello from a killable subprocess")
    out = tmp_path / "t.docx"
    ok, err = pdf_import._pdf_to_docx(src, out)
    assert ok is True, err
    assert out.exists() and out.stat().st_size > 0


def test_pdf_to_docx_subprocess_rejects_garbage(tmp_path):
    """Un PDF illeggibile → il figlio esce non-zero, _pdf_to_docx ritorna un
    errore pulito e rimuove l'eventuale output parziale."""
    pytest.importorskip("pdf2docx")
    src = tmp_path / "bad.pdf"
    src.write_bytes(b"not a pdf at all")
    out = tmp_path / "bad.docx"
    ok, err = pdf_import._pdf_to_docx(src, out)
    assert ok is False
    assert not out.exists()
    assert "PDF import" in err  # messaggio pulito, mai uno stacktrace grezzo


@pytest.mark.skipif(not pdf_import._ocr_available(),
                    reason="OCR chain (ocrmypdf+ghostscript+tesseract) not installed")
def test_ocr_pdf_to_docx_from_real_invisible_layer(tmp_path):
    """IL caso che ha motivato il fix: una scansione passa da OCRmyPDF (layer di
    testo INVISIBILE, alpha 0) e il docx risultante NON deve essere vuoto —
    pdf2docx lo produrrebbe vuoto (bug empirico che _ocr_pdf_to_docx corregge)."""
    from docx import Document
    from PIL import ImageDraw, ImageFont

    img = Image.new("RGB", (1200, 400), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype("arial.ttf", 96)
    draw.text((100, 120), "HELLO CONVERTER OCR", fill="black", font=font)
    scan = tmp_path / "scan.pdf"
    img.save(scan, "PDF", resolution=150.0)

    ocred = tmp_path / "scan.ocr.pdf"
    ok, err = pdf_import._ocr_pdf(scan, ocred)
    assert ok is True, err

    out = tmp_path / "o.docx"
    ok, err = pdf_import._ocr_pdf_to_docx(ocred, out)
    assert ok is True, err
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "".join(text.split()) != "", "il docx dal PDF OCR-izzato NON deve essere vuoto"
    assert "CONVERTER" in text.upper()
