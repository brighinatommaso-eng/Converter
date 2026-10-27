"""Motore markup in-process (ha sostituito pandoc): lettori, scrittori, verdetti.

Le conversioni passano da run_job come nella GUI; dove serve si apre il file
prodotto con la libreria del formato (python-docx, zipfile, defusedxml) o con
LibreOffice (skip se assente) per verificare che sia davvero leggibile.
"""

from __future__ import annotations

import subprocess
import zipfile
from pathlib import Path

import pytest

import converter.converters  # noqa: F401  popola il registry
from converter.converters.markupdoc.readers import markdown_to_html
from converter.converters.markupdoc.tree import parse_html
from converter.core import deps, engine
from converter.core.job import ConversionJob
from converter.core.pipeline import run_job
from converter.core.registry import resolve

soffice = pytest.mark.skipif(deps.find_soffice() is None, reason="LibreOffice not installed")

_SAMPLE = """# Report *Q3*

Intro with **bold**, *italic*, ~~gone~~, `x = 1` and [a link](https://example.com).
Line one
second line with a hard break\\
after it.

> Quoted text

1. first
2. second
   - nested a
   - nested b
3. third

```python
def f():
    return "`tick`"
```

| Name | Value |
|:-----|------:|
| a | 1 |
| b \\| c | 2 |

---

Last paragraph.
"""


@pytest.fixture(autouse=True)
def _portable():
    engine.set_mode("PORTABLE")       # nessun LibreOffice: tutto in-process
    yield
    engine.set_mode("FULL")


def _convert(src: Path, src_fmt: str, dst_fmt: str, out: Path):
    result = run_job(ConversionJob(src, src_fmt, dst_fmt, out))
    assert result.ok, result.error
    return result


def _md(tmp_path: Path, text: str = _SAMPLE, name: str = "doc.md") -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def _png(path: Path, size=(40, 20), color="red") -> Path:
    from PIL import Image

    Image.new("RGB", size, color).save(path)
    return path


# --- Albero e ripulitura ------------------------------------------------------------

def test_untrusted_html_loses_scripts_events_and_javascript_links():
    body, title, _lang = parse_html(
        '<html><head><title>T</title><script>alert(1)</script><style>p{}</style></head>'
        '<body><p onclick="steal()">ok <a href="javascript:alert(1)">x</a>'
        '<a href="data:text/html,<script>">y</a><a href="https://ok.example">z</a></p>'
        '<iframe src="https://evil"></iframe></body></html>')
    tags = [node.tag for node in body.iter()]
    assert title == "T"
    assert "script" not in tags and "iframe" not in tags
    links = [node.attrs for node in body.iter() if node.tag == "a"]
    assert links == [{}, {}, {"href": "https://ok.example"}]
    assert all("onclick" not in node.attrs for node in body.iter())
    assert "alert" not in body.text()


def test_dirty_html_is_repaired_like_a_browser():
    body, _t, _l = parse_html("<div>Hi <b>there</b><p>one<li>stray</div>"
                              "<p>before<ul><li>item</ul>after</p>")
    blocks = [(node.tag, " ".join(node.text().split())) for node in body.children]
    assert blocks == [("p", "Hi there"), ("p", "one"), ("p", "stray"), ("p", "before"),
                      ("ul", "item"), ("p", "after")]


def test_html_charset_from_meta_is_honoured(tmp_path):
    src = tmp_path / "page.html"
    src.write_bytes('<html><head><meta charset="windows-1252"></head><body><p>caffè €5</p>'
                    "</body></html>".encode("cp1252"))
    out = tmp_path / "page.txt"
    _convert(src, "html", "txt", out)
    assert "caffè €5" in out.read_text(encoding="utf-8")


# --- Markdown ----------------------------------------------------------------------

def test_md_to_html_is_lossless_and_uses_reference_rendering(tmp_path):
    out = tmp_path / "doc.html"
    result = _convert(_md(tmp_path), "md", "html", out)
    assert result.verdict.is_lossless
    html = out.read_text(encoding="utf-8")
    for fragment in ("<h1>Report <em>Q3</em></h1>", "<strong>bold</strong>", "<s>gone</s>",
                     "<code>x = 1</code>", '<a href="https://example.com">a link</a>',
                     "<blockquote>", '<code class="language-python">',
                     'style="text-align:right"', "<hr>", "<br>"):
        assert fragment in html, fragment


def test_raw_html_in_markdown_makes_the_verdict_lossy(tmp_path):
    out = tmp_path / "doc.html"
    result = _convert(_md(tmp_path, "Hi <span style='color:red' onclick='x'>red</span>\n"),
                      "md", "html", out)
    assert result.verdict.is_lossy and "Raw HTML" in result.verdict.reason
    assert "onclick" not in out.read_text(encoding="utf-8")


def test_md_round_trip_through_html_keeps_text_and_structure(tmp_path):
    """md → html → md → html: il secondo rendering è identico al primo (nessun
    simbolo letto come markup, nessuna struttura persa)."""
    html1 = tmp_path / "a.html"
    _convert(_md(tmp_path), "md", "html", html1)
    md2 = tmp_path / "b.md"
    _convert(html1, "html", "md", md2)
    first, _t, _l = parse_html(markdown_to_html(_SAMPLE))
    second, _t, _l = parse_html(markdown_to_html(md2.read_text(encoding="utf-8")))

    def shape(node):
        return [(c.tag, " ".join(c.text().split())) for c in node.iter() if c.tag != "body"]

    assert shape(second) == shape(first)


def test_markdown_writer_escapes_literal_symbols(tmp_path):
    src = tmp_path / "a.html"
    src.write_text("<p># not a heading *not em* [x](y) a|b &amp; &lt;tag&gt; 1. no</p>"
                   "<p>2) no list</p>", encoding="utf-8")
    out = tmp_path / "a.md"
    _convert(src, "html", "md", out)
    back, _t, _l = parse_html(markdown_to_html(out.read_text(encoding="utf-8")))
    assert [c.tag for c in back.children] == ["p", "p"]
    assert back.children[0].text() == "# not a heading *not em* [x](y) a|b & <tag> 1. no"


def test_markdown_lists_keep_start_numbers_and_code_spans_with_backticks(tmp_path):
    src = tmp_path / "a.html"
    src.write_text('<ol start="3"><li>three</li><li>four<ol><li>sub</li></ol></li></ol>'
                   "<p><code>a ` b</code></p>", encoding="utf-8")
    out = tmp_path / "a.md"
    _convert(src, "html", "md", out)
    text = out.read_text(encoding="utf-8")
    assert "3. three\n4. four\n   1. sub" in text
    assert "``a ` b``" in text


def test_merged_cells_in_markdown_are_declared(tmp_path):
    src = tmp_path / "t.html"
    src.write_text('<table><tr><th colspan="2">wide</th></tr><tr><td>a</td><td>b</td></tr></table>',
                   encoding="utf-8")
    out = tmp_path / "t.md"
    result = _convert(src, "html", "md", out)
    assert "Merged table cells" in result.verdict.reason
    assert "| wide |  |" in out.read_text(encoding="utf-8")


# --- Immagini: locali incorporate, remote mai scaricate ------------------------------

def test_local_image_is_embedded_and_remote_one_is_declared(tmp_path):
    _png(tmp_path / "pic.png")
    src = _md(tmp_path, "![local](pic.png)\n\n![remote](https://example.com/x.png)\n")
    html = tmp_path / "a.html"
    result = _convert(src, "md", "html", html)
    text = html.read_text(encoding="utf-8")
    assert "data:image/png;base64," in text            # l'HTML resta un file solo
    assert 'src="https://example.com/x.png"' in text    # il remoto resta un link
    assert result.verdict.is_lossless

    docx = tmp_path / "a.docx"
    result = _convert(src, "md", "docx", docx)
    from docx import Document

    assert len(Document(str(docx)).inline_shapes) == 1
    assert result.verdict.is_lossy and "offline" in result.verdict.reason


def test_missing_local_image_is_declared(tmp_path):
    src = _md(tmp_path, "![gone](nope.png)\n\ntext\n")
    result = _convert(src, "md", "docx", tmp_path / "a.docx")
    assert "could not be found" in result.verdict.reason


def test_non_image_files_are_never_embedded(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("password=hunter2", encoding="utf-8")
    src = _md(tmp_path, "![x](secret.txt)\n")
    out = tmp_path / "a.html"
    _convert(src, "md", "html", out)
    assert "aHVudGVy" not in out.read_text(encoding="utf-8")   # base64 di "hunter"


# --- DOCX ----------------------------------------------------------------------------

def test_md_to_docx_uses_native_word_structure(tmp_path):
    from docx import Document

    out = tmp_path / "a.docx"
    _convert(_md(tmp_path, _SAMPLE + "\n1. again\n2. twice\n"), "md", "docx", out)
    doc = Document(str(out))
    styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]
    assert ("Heading 1", "Report Q3") in styles
    assert ("Quote", "Quoted text") in styles
    assert ("List Number", "first") in styles and ("List Bullet 2", "nested a") in styles
    xml = doc.element.xml
    assert "w:hyperlink" in xml and "Consolas" in xml and "w:strike" in xml
    assert "w:pBdr" in xml                                     # la linea orizzontale
    # Il secondo elenco numerato riparte da 1: numId diverso dal primo.
    num_ids = []
    for p in doc.paragraphs:
        if p.style.name == "List Number":
            num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
            num_ids.append(num_pr.numId.val if num_pr is not None else None)
    assert len(set(num_ids)) == 2
    table = doc.tables[0]
    assert [c.text for c in table.rows[2].cells] == ["b | c", "2"]


def test_docx_table_keeps_merged_cells(tmp_path):
    from docx import Document

    src = tmp_path / "t.html"
    src.write_text('<table><tr><td colspan="2">wide</td></tr><tr><td rowspan="2">tall</td>'
                   "<td>x</td></tr><tr><td>y</td></tr></table>", encoding="utf-8")
    out = tmp_path / "t.docx"
    _convert(src, "html", "docx", out)
    xml = Document(str(out)).tables[0]._tbl.xml
    assert 'w:gridSpan w:val="2"' in xml and "w:vMerge" in xml


# --- ODT -----------------------------------------------------------------------------

def _odt_content(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("content.xml").decode("utf-8")


def test_md_to_odt_has_real_structure(tmp_path):
    out = tmp_path / "a.odt"
    _convert(_md(tmp_path), "md", "odt", out)
    content = _odt_content(out)
    for fragment in ('text:outline-level="1"', "<text:list", "<table:table ", 'xlink:href="https://example.com"',
                     "<text:line-break/>"):
        assert fragment in content, fragment
    with zipfile.ZipFile(out) as archive:
        styles = archive.read("styles.xml").decode("utf-8")
    assert 'style:name="Heading_20_1"' in styles and 'style:name="Quotations"' in styles


def test_odt_round_trip_keeps_code_indentation_and_structure(tmp_path):
    odt = tmp_path / "a.odt"
    _convert(_md(tmp_path), "md", "odt", odt)
    back = tmp_path / "back.md"
    _convert(odt, "odt", "md", back)
    text = back.read_text(encoding="utf-8")
    assert '    return "`tick`"' in text               # 4 spazi, non 3
    assert "# Report *Q3*" in text and "1. first" in text and "   - nested a" in text
    assert "---" in text and "| b \\| c | 2 |" in text


def test_odt_reader_handles_styles_lists_tables_and_footnotes(tmp_path):
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import H, List, ListItem, Note, NoteBody, NoteCitation, P, Span

    doc = OpenDocumentText()
    bold = Style(name="B", family="text")
    bold.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(bold)
    doc.text.addElement(H(outlinelevel="2", text="Chapter"))
    para = P(text="Plain ")
    para.addElement(Span(stylename=bold, text="strong"))
    note = Note(noteclass="footnote")
    note.addElement(NoteCitation(text="1"))
    body = NoteBody()
    body.addElement(P(text="The footnote."))
    note.addElement(body)
    para.addElement(note)
    doc.text.addElement(para)
    lst = List()
    item = ListItem()
    item.addElement(P(text="bullet"))
    lst.addElement(item)
    doc.text.addElement(lst)
    table = Table(name="T")
    table.addElement(TableColumn(numbercolumnsrepeated="2"))
    row = TableRow()
    for value in ("c1", "c2"):
        cell = TableCell()
        cell.addElement(P(text=value))
        row.addElement(cell)
    table.addElement(row)
    doc.text.addElement(table)
    src = tmp_path / "in.odt"
    doc.save(str(src))

    out = tmp_path / "out.md"
    result = _convert(src, "odt", "md", out)
    text = out.read_text(encoding="utf-8")
    assert "## Chapter" in text and "Plain **strong**<sup>1</sup>" in text
    assert "- bullet" in text and "| c1 | c2 |" in text
    assert "1. The footnote." in text
    assert "Footnotes were moved" in result.verdict.reason


@soffice
def test_libreoffice_opens_the_odt_and_rtf_we_write(tmp_path):
    for fmt in ("odt", "rtf"):
        out = tmp_path / f"a.{fmt}"
        _convert(_md(tmp_path), "md", fmt, out)
        subprocess.run([deps.find_soffice(), "--headless", "--convert-to", "docx", "--outdir",
                        str(tmp_path / fmt), str(out)], capture_output=True, timeout=180,
                       env=deps.clean_subprocess_env())
        from docx import Document

        converted = Document(str(tmp_path / fmt / "a.docx"))
        styles = [(p.style.name, p.text) for p in converted.paragraphs if p.text.strip()]
        assert ("Heading 1", "Report Q3") in styles, (fmt, styles[:3])
        assert any("Last paragraph." in text for _s, text in styles), fmt


# --- RTF -----------------------------------------------------------------------------

def test_rtf_escapes_control_characters_and_unicode(tmp_path):
    src = tmp_path / "a.html"
    src.write_text("<h2>{braces} \\back</h2><p>àé € 😀</p>", encoding="utf-8")
    out = tmp_path / "a.rtf"
    _convert(src, "html", "rtf", out)
    rtf = out.read_text(encoding="ascii")
    assert rtf.startswith("{\\rtf1") and rtf.count("{") == rtf.count("}") + rtf.count("\\{") - rtf.count("\\}")
    assert "\\{braces\\} \\\\back" in rtf and "\\outlinelevel1" in rtf
    assert "\\u224?" in rtf and "\\u8364?" in rtf and "\\u-10179?\\u-8704?" in rtf


# --- EPUB ----------------------------------------------------------------------------

def test_md_to_epub_is_a_valid_package_and_reads_back(tmp_path):
    from defusedxml import ElementTree

    _png(tmp_path / "pic.png")
    src = _md(tmp_path, "# One\n\nText & more ![p](pic.png)\n\n## Sub\n\nx\n\n# Two\n\ny\n")
    epub = tmp_path / "a.epub"
    _convert(src, "md", "epub", epub)
    with zipfile.ZipFile(epub) as archive:
        first = archive.infolist()[0]
        assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
        assert archive.read("mimetype") == b"application/epub+zip"
        names = archive.namelist()
        assert {"OEBPS/chapter-001.xhtml", "OEBPS/chapter-002.xhtml", "OEBPS/images/img-001.png",
                "OEBPS/nav.xhtml", "OEBPS/toc.ncx"} <= set(names)
        for name in names:
            if name.endswith((".xhtml", ".opf", ".ncx", ".xml")):
                ElementTree.fromstring(archive.read(name))      # ben formato
        nav = archive.read("OEBPS/nav.xhtml").decode("utf-8")
    assert "One" in nav and "Sub" in nav and "Two" in nav
    back = tmp_path / "back.md"
    _convert(epub, "epub", "md", back)
    text = back.read_text(encoding="utf-8")
    assert "# One" in text and "## Sub" in text and "# Two" in text and "data:image/png" in text


def test_drm_epub_is_refused_with_a_clear_message(tmp_path):
    epub = tmp_path / "drm.epub"
    with zipfile.ZipFile(epub, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/encryption.xml",
                         '<encryption><EncryptedData><CipherData>'
                         '<CipherReference URI="OEBPS/chapter.xhtml"/></CipherData>'
                         "</EncryptedData></encryption>")
    result = run_job(ConversionJob(epub, "epub", "md", tmp_path / "o.md"))
    assert not result.ok and "DRM" in result.error


def test_archive_bomb_is_refused(tmp_path, monkeypatch):
    import converter.converters.markupdoc.readers as readers

    monkeypatch.setattr(readers, "_MAX_ARCHIVE_BYTES", 1000)
    epub = tmp_path / "big.epub"
    with zipfile.ZipFile(epub, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("filler", "0" * 100_000)
    result = run_job(ConversionJob(epub, "epub", "md", tmp_path / "o.md"))
    assert not result.ok and "unsafe" in result.error


def test_not_an_epub_gives_a_readable_error(tmp_path):
    epub = tmp_path / "fake.epub"
    epub.write_text("not a zip", encoding="utf-8")
    result = run_job(ConversionJob(epub, "epub", "html", tmp_path / "o.html"))
    assert not result.ok and "ZIP" in result.error
    assert not (tmp_path / "o.html").exists()


# --- Testo semplice ----------------------------------------------------------------

def test_txt_to_html_and_docx_keep_every_space_and_line(tmp_path):
    from docx import Document

    src = tmp_path / "a.txt"
    src.write_text("col1    col2\n  indented\n\nsecond <p> & para\n", encoding="utf-8")
    html = tmp_path / "a.html"
    result = _convert(src, "txt", "html", html)
    assert result.verdict.is_lossless
    text = html.read_text(encoding="utf-8")
    assert "white-space:pre-wrap" in text and "col1    col2<br>  indented" in text
    assert "&lt;p&gt; &amp; para" in text
    docx = tmp_path / "a.docx"
    _convert(src, "txt", "docx", docx)
    paragraphs = [p.text for p in Document(str(docx)).paragraphs if p.text]
    assert paragraphs == ["col1    col2\n  indented", "second <p> & para"]


def test_empty_document_is_an_error_not_an_empty_file(tmp_path):
    src = tmp_path / "empty.html"
    src.write_text("<html><body><script>x</script></body></html>", encoding="utf-8")
    result = run_job(ConversionJob(src, "html", "docx", tmp_path / "o.docx"))
    assert not result.ok and "no readable content" in result.error


# --- Instradamento -------------------------------------------------------------------

@pytest.mark.parametrize("mode", ["FULL", "PORTABLE"])
def test_markup_pairs_are_routed_in_process_in_both_modes(mode, monkeypatch):
    monkeypatch.setattr(deps, "find_soffice", lambda: "/fake/soffice")
    engine.set_mode(mode)
    for src, dst in (("md", "docx"), ("md", "odt"), ("md", "rtf"), ("md", "epub"), ("html", "docx"),
                     ("epub", "md"), ("txt", "docx"), ("odt", "md"), ("odt", "txt")):
        assert type(resolve(src, dst)).__name__ == "MarkupConverter", (src, dst)
    assert type(resolve("txt", "md")).__name__ == "PlainTextMarkdownConverter"
    # odt verso formati office resta a LibreOffice (layout), come html/txt → pdf.
    assert type(resolve("odt", "docx")).__name__ == "DocumentConverter"
    assert type(resolve("html", "pdf")).__name__ == "DocumentConverter"
