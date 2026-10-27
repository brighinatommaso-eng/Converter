"""Analisi del layout con il modello AI (docmodel.layout_ai + _ai_page_blocks).

Due livelli:
  * l'assemblaggio guidato dalle regioni, con regioni finte: non serve il modello,
    gira ovunque e fissa le regole (ordine, niente testo perso, tabelle, titoli);
  * end-to-end con il modello vero su un documento generato da LibreOffice: gira
    solo se modello e soffice sono installati (in CI il job Linux li ha).
"""

from __future__ import annotations

import shutil
import subprocess
import threading
from pathlib import Path

import pytest

from converter.converters.docmodel import layout_ai
from converter.converters.docmodel.layout_ai import Region
from converter.converters.docmodel.pdfreader import (
    ReadCancelled,
    _ai_page_blocks,
    _starts_item,
    _table_from_region_lines,
    _would_fit,
)
from converter.converters.docmodel.types import (
    ImageBlock,
    Line,
    ListBlock,
    Paragraph,
    Run,
    Table,
)

pytestmark = pytest.mark.layout_ai


def _line(text: str, x: float, y: float, *, size: float = 10.0, bold: bool = False,
          width: float | None = None) -> Line:
    width = width if width is not None else 5.0 * len(text)
    words, cursor = [], x
    for word in text.split():
        words.append((cursor, cursor + 5.0 * len(word), word))
        cursor += 5.0 * (len(word) + 1)
    return Line(runs=[Run(text=text, size=size, bold=bold)], x=x, y=y, width=width,
                height=size, words=words)


def _page(lines, regions, images=()):
    return {"lines": lines, "segments": [], "images": list(images), "regions": regions,
            "index": 0, "width": 595.0, "height": 842.0}


def _texts(blocks):
    out = []
    for block in blocks:
        if isinstance(block, Paragraph):
            out.append(block.text)
        elif isinstance(block, ListBlock):
            out += [item.paragraph.text for item in block.items]
    return out


# --- Assemblaggio con regioni finte ----------------------------------------

def test_reading_order_follows_the_model_not_the_page_height():
    """Due colonne: la sinistra si legge tutta prima della destra, anche se la
    prima riga della destra sta più in alto dell'ultima della sinistra."""
    left = [_line("left one", 60, 700), _line("left two", 60, 688)]
    right = [_line("right one", 320, 700), _line("right two", 320, 688)]
    regions = [Region("text", 0.9, 55, 680, 290, 712, 0),
               Region("text", 0.9, 315, 680, 560, 712, 1)]
    blocks = _ai_page_blocks(_page(left + right, regions), left + right, [], [], None, 60, 560)
    assert _texts(blocks) == ["left one left two", "right one right two"]


def test_text_outside_every_region_is_kept():
    """Il modello può solo riordinare: una riga che non cade in nessuna regione
    resta nel documento, vicino alla regione che la precede."""
    lines = [_line("inside the region", 60, 700), _line("forgotten line", 60, 600)]
    regions = [Region("text", 0.9, 55, 690, 300, 712, 0)]
    blocks = _ai_page_blocks(_page(lines, regions), lines, [], [], None, 60, 560)
    assert _texts(blocks) == ["inside the region", "forgotten line"]


def test_title_regions_become_headings_and_page_numbers_are_dropped():
    from converter.converters.docmodel.pdfreader import _classify_headings
    from converter.converters.docmodel.types import Document

    lines = [_line("Results", 60, 720), _line("body text " * 8, 60, 700, width=400),
             _line("12", 290, 40)]
    regions = [Region("paragraph_title", 0.9, 55, 715, 120, 732, 0),
               Region("text", 0.9, 55, 695, 470, 712, 1),
               Region("number", 0.9, 285, 35, 305, 52, 2)]
    blocks = _ai_page_blocks(_page(lines, regions), lines, [], [], None, 60, 470)
    document = Document(blocks=blocks)
    _classify_headings(document)
    assert [b.text for b in document.blocks] == ["Results", ("body text " * 8).strip()]
    assert document.blocks[0].style == "heading"     # stesso corpo, ma il modello dice titolo
    assert document.blocks[1].style == "body"


def test_figure_region_keeps_the_original_photo_and_drops_its_labels():
    photo = ImageBlock(data=b"\x89PNG fake", ext="png", width_pt=200, height_pt=100, x=60, y=500)
    label = _line("axis label", 70, 520)
    regions = [Region("image", 0.9, 58, 498, 262, 602, 0)]
    blocks = _ai_page_blocks(_page([label], regions, [photo]), [label], [], [], None, 60, 560)
    assert blocks == [photo]


def test_full_page_background_goes_after_the_text():
    background = ImageBlock(data=b"x", ext="jpg", width_pt=595, height_pt=842, x=0, y=0)
    text = _line("headline", 60, 700)
    regions = [Region("text", 0.9, 55, 690, 300, 712, 0)]
    blocks = _ai_page_blocks(_page([text], regions, [background]), [text], [], [], None, 60, 560)
    assert isinstance(blocks[0], Paragraph) and blocks[-1] is background


def test_lines_split_by_the_model_rejoin_into_one_paragraph():
    """Il modello a volte vede ogni riga di un paragrafo come una regione."""
    lines = [_line("a long first line that fills the column", 60, 700, width=300),
             _line("and its continuation", 60, 688, width=120)]
    regions = [Region("text", 0.9, 55, 698, 365, 712, 0),
               Region("text", 0.9, 55, 686, 365, 699, 1)]
    blocks = _ai_page_blocks(_page(lines, regions), lines, [], [], None, 60, 360)
    assert _texts(blocks) == ["a long first line that fills the column and its continuation"]


def test_justified_line_split_by_pdfminer_is_one_line_again():
    pieces = [_line("facto method.", 60, 700), _line("In this work, we show", 160, 700)]
    regions = [Region("text", 0.9, 55, 695, 300, 712, 0)]
    blocks = _ai_page_blocks(_page(pieces, regions), pieces, [], [], None, 60, 300)
    assert _texts(blocks) == ["facto method. In this work, we show"]


def test_hard_line_breaks_stay_separate():
    """Righe corte andate a capo apposta (la parola dopo ci stava): non si fondono."""
    prev = _line("For an inclusive ecosystem", 60, 700)
    nxt = _line("For digital transformation", 60, 688)
    assert _would_fit(prev, nxt, 300.0)
    wrapped = _line("a line that really reaches the margin", 60, 700, width=238)
    assert not _would_fit(wrapped, _line("continuation", 60, 688), 300.0)


def test_middle_dot_is_not_a_bullet_inside_a_sentence():
    assert not _starts_item("· Character Recognition · Toolkit", "Keywords: Deep Learning")
    assert _starts_item("· second item", "· first item")
    assert _starts_item("• real bullet", "Some text")


def test_table_region_rebuilds_rows_with_multiline_cells_and_row_labels():
    lines = [
        _line("Store", 60, 700, bold=True), _line("Partners", 160, 700, bold=True),
        _line("Coffee", 60, 686), _line("Starbucks, Twosome,", 160, 686),
        _line("Pascucci, Ediya", 160, 674),
        _line("NLL", 160, 640), _line("DP", 60, 634), _line("Triplet", 160, 628),
    ]
    table = _table_from_region_lines(lines, bordered=True)
    assert isinstance(table, Table) and table.header_rows == 1
    assert table.text_grid() == [["Store", "Partners"],
                                 ["Coffee", "Starbucks, Twosome, Pascucci, Ediya"],
                                 ["DP", "NLL Triplet"]]


# --- Lettore PDF senza modello ----------------------------------------------

def _raw_pdf(path: Path, page_content: bytes, extra: list[bytes] = ()) -> None:
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> /XObject << /X1 6 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(page_content) + page_content + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", *extra]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def test_text_inside_a_form_xobject_is_read(tmp_path):
    """Regressione: il testo nei Form XObject (impaginati InDesign, report, moduli)
    spariva in silenzio, perché pdfminer non lo raggruppa in righe senza all_texts."""
    from converter.converters.docmodel import read_pdf

    form = b"BT /F1 12 Tf 72 700 Td (Text inside a form XObject) Tj ET"
    xobject = (b"<< /Type /XObject /Subtype /Form /BBox [0 0 595 842] "
               b"/Resources << /Font << /F1 5 0 R >> >> /Length %d >>\nstream\n" % len(form)
               + form + b"\nendstream")
    pdf = tmp_path / "form.pdf"
    _raw_pdf(pdf, b"BT /F1 12 Tf 72 760 Td (Page text) Tj ET\n/X1 Do", [xobject])
    texts = [b.text for b in read_pdf(pdf, use_ai=False).blocks if isinstance(b, Paragraph)]
    assert "Text inside a form XObject" in texts and "Page text" in texts


def test_without_the_model_the_verdict_says_so(tmp_path, monkeypatch):
    from converter.converters import textdoc
    from converter.core.job import CancelToken, ConversionJob
    from converter.core.pipeline import run_job

    monkeypatch.setenv("CONVERTER_LAYOUT_AI", "0")
    pdf = tmp_path / "plain.pdf"
    _raw_pdf(pdf, b"BT /F1 12 Tf 72 760 Td (Hello) Tj ET")
    assert textdoc  # registrazione del converter
    result = run_job(ConversionJob(src_path=pdf, src_format="pdf", dst_format="md",
                                   output_path=tmp_path / "out.md", cancel=CancelToken()))
    assert result.ok
    assert "AI layout analysis was not used" in result.verdict.reason


def test_cancel_stops_the_reader(tmp_path):
    from converter.converters.docmodel import read_pdf

    pdf = tmp_path / "plain.pdf"
    _raw_pdf(pdf, b"BT /F1 12 Tf 72 760 Td (Hello) Tj ET")
    with pytest.raises(ReadCancelled):
        read_pdf(pdf, use_ai=False, cancel=lambda: True)


# --- End-to-end con il modello vero ------------------------------------------

_needs_model = pytest.mark.skipif(
    layout_ai.unavailable_reason() not in (None, "AI layout analysis is turned off."),
    reason="AI layout model not installed")
_needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")


@pytest.fixture(scope="module")
def two_column_pdf(tmp_path_factory) -> Path:
    """Articolo a due colonne con titolo, tabella con bordi e un grafico, via LibreOffice."""
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from PIL import Image, ImageDraw

    folder = tmp_path_factory.mktemp("ai")
    chart = folder / "chart.png"
    image = Image.new("RGB", (600, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.line([(40, 260), (580, 260)], fill="black", width=3)
    draw.line([(40, 20), (40, 260)], fill="black", width=3)
    draw.line([(40, 250), (200, 150), (360, 180), (560, 40)], fill="blue", width=6)
    image.save(chart)

    document = docx.Document()
    document.add_heading("Layout Analysis Benchmark Report", level=0)
    section = document.sections[0]
    cols = section._sectPr.makeelement(qn("w:cols"), {qn("w:num"): "2", qn("w:space"): "720"})
    section._sectPr.append(cols)
    document.add_heading("1 Introduction", level=1)
    for index in range(1, 5):
        paragraph = document.add_paragraph(
            f"Paragraph {index} of the left column explains the method in plain words, "
            f"with enough sentences to wrap over several lines of the narrow column. " * 3)
        paragraph.runs[0].font.size = Pt(10)
    document.add_heading("2 Results", level=1)
    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    for grid_col in table._tbl.tblGrid.findall(qn("w:gridCol")):
        grid_col.set(qn("w:w"), str(70 * 20))          # 70pt: sta nella colonna
    for column in table.columns:
        for cell in column.cells:
            cell.width = Pt(70)
    for row, values in enumerate([("Model", "Score", "Time"), ("Alpha", "91.2", "1.3"),
                                  ("Beta", "88.7", "0.9"), ("Gamma", "93.5", "2.1")]):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    document.add_paragraph("Table 1: scores of the three models.")
    document.add_picture(str(chart), width=Pt(220))
    document.add_paragraph("Figure 1: score over time.")
    for index in range(5, 8):
        document.add_paragraph(
            f"Paragraph {index} discusses the results shown above and closes the "
            f"report with a short conclusion about the measured numbers. " * 3)
    source = folder / "report.docx"
    document.save(source)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(folder),
                    str(source)], capture_output=True, timeout=180, check=True)
    return folder / "report.pdf"


@_needs_model
@_needs_soffice
def test_model_rebuilds_a_two_column_report(two_column_pdf):
    from converter.converters.docmodel import read_pdf
    from converter.converters.docmodel.pdfreader import _classify_headings  # noqa: F401

    document = read_pdf(two_column_pdf)
    assert document.ai_pages >= 1 and not document.ai_unavailable
    blocks = [b for b in document.blocks if not type(b).__name__ == "PageBreak"]
    texts = _texts(blocks)
    joined = " ".join(texts)
    # Ordine di lettura: tutti i paragrafi 1..7 nell'ordine giusto.
    positions = [joined.find(f"Paragraph {i} ") for i in range(1, 8)]
    assert all(p >= 0 for p in positions) and positions == sorted(positions)
    headings = [b.text for b in blocks if isinstance(b, Paragraph) and b.style in ("heading", "title")]
    assert "1 Introduction" in headings and "2 Results" in headings
    tables = [b for b in blocks if isinstance(b, Table)]
    assert tables and tables[0].text_grid()[1] == ["Alpha", "91.2", "1.3"]
    assert any(isinstance(b, ImageBlock) for b in blocks)      # il grafico c'è


@_needs_model
@_needs_soffice
def test_parallel_conversions_share_the_model_safely(two_column_pdf):
    """La GUI converte in parallelo: PDFium non è thread-safe, il lock lo protegge."""
    from converter.converters.docmodel import read_pdf

    results, errors = [], []

    def work():
        try:
            results.append(read_pdf(two_column_pdf).ai_pages)
        except Exception as exc:   # pragma: no cover - il test fallisce sotto
            errors.append(exc)

    threads = [threading.Thread(target=work) for _ in range(3)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert not errors and len(results) == 3 and len(set(results)) == 1
